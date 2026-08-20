from io import BytesIO
from urllib.parse import urlencode

from django.contrib.auth.decorators import login_required
from django.contrib import messages
from django.core.paginator import Paginator
from django.db.models import Q
from django.http import HttpResponse
from django.shortcuts import render, redirect, get_object_or_404
from django.utils import timezone

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.units import cm
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle

from alertes_et_notifications.models import Alerte
from .models import ObjetOublie
from entreprise.models import Departement
from utilisateurs.models import HistoriqueAction
from visites.models import Visite, Visiteur


def _base_ctx():
    return {
        'departements': Departement.objects.all(),
        'visites': Visite.objects.select_related('visiteur')[:100],
        'visiteurs': Visiteur.objects.filter(statut='ACTIF').order_by('nom', 'prenom'),
    }


def _export_pdf(rows, headers, title, filename, col_ratios=None):
    buffer = BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=landscape(A4), title=title,
                            leftMargin=1.5*cm, rightMargin=1.5*cm)
    styles = getSampleStyleSheet()
    cell_style = ParagraphStyle('CellStyle', parent=styles['Normal'],
                                fontSize=8, leading=10, alignment=1)
    elements = []
    elements.append(Paragraph(title, styles['Title']))
    elements.append(Spacer(1, 0.5 * cm))
    elements.append(Paragraph(f"Généré le {timezone.now().strftime('%d/%m/%Y %H:%M')}", styles['Normal']))
    elements.append(Spacer(1, 0.5 * cm))
    col_count = len(headers)
    available_width = landscape(A4)[0] - 3 * cm
    if col_ratios is None:
        col_ratios = [1] * col_count
    total_ratio = sum(col_ratios)
    col_widths = [available_width * r / total_ratio for r in col_ratios]
    wrapped_rows = [[Paragraph(str(c), cell_style) for c in row] for row in rows]
    data = [headers] + wrapped_rows
    table = Table(data, colWidths=col_widths, repeatRows=1)
    table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1270b8')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, -1), 8),
        ('BOTTOMPADDING', (0, 0), (-1, 0), 8),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#e2e8f0')),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.HexColor('#f8fafc'), colors.white]),
        ('VALIGN', (0, 0), (-1, -1), 'TOP'),
    ]))
    elements.append(table)
    doc.build(elements)
    pdf = buffer.getvalue()
    buffer.close()
    response = HttpResponse(pdf, content_type='application/pdf')
    response['Content-Disposition'] = f'inline; filename="{filename}_{timezone.now().strftime("%Y%m%d")}.pdf"'
    return response


def _export_excel(rows, headers, title, filename):
    wb = Workbook()
    ws = wb.active
    ws.title = title[:31]
    ws.append(headers)
    for row in rows:
        ws.append(row)
    for col_idx in range(1, len(headers) + 1):
        ws.column_dimensions[chr(64 + col_idx)].width = 22
    header_fill = PatternFill(start_color='1270b8', end_color='1270b8', fill_type='solid')
    header_font = Font(bold=True, color='ffffff', size=11)
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    response = HttpResponse(buffer, content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = f'attachment; filename="{filename}_{timezone.now().strftime("%Y%m%d")}.xlsx"'
    return response


@login_required
def liste_objets_oublies(request):
    query = request.GET.get('q', '').strip()
    statut = request.GET.get('statut', '')
    items = ObjetOublie.objects.select_related(
        'visite__visiteur', 'departement'
    )
    if query:
        items = items.filter(
            Q(nom_objet__icontains=query) | Q(lieu_trouve__icontains=query) |
            Q(description__icontains=query) | Q(categorie__icontains=query) |
            Q(personne_recupere__icontains=query) |
            Q(visite__visiteur__nom__icontains=query) | Q(visite__visiteur__prenom__icontains=query)
        )
    if statut and statut != 'TOUS':
        items = items.filter(statut=statut)
    items = items.order_by('-created_at')
    paginator = Paginator(items, 10)
    page_obj = paginator.get_page(request.GET.get('page', 1))
    page_links = paginator.get_elided_page_range(page_obj.number, on_each_side=2, on_ends=1)
    fp = {k: v for k, v in [('q', query), ('statut', statut)] if v and v != 'TOUS'}
    return render(request, 'objets_oublies/liste.html', {
        'page_obj': page_obj, 'page_links': page_links, 'query': query,
        'statut_filter': statut,
        'filter_params': urlencode(fp) + '&' if fp else '',
    })


@login_required
def ajouter_objet_oublie(request):
    if request.method == 'POST':
        nom_objet = request.POST.get('nom_objet', '').strip()
        if not nom_objet:
            messages.error(request, 'Le nom de l\'objet est requis.')
            return render(request, 'objets_oublies/ajouter.html', _base_ctx())
        obj = ObjetOublie.objects.create(
            nom_objet=nom_objet,
            categorie=request.POST.get('categorie', '').strip() or None,
            description=request.POST.get('description', '').strip() or None,
            date_trouve=request.POST.get('date_trouve', '').strip() or None,
            lieu_trouve=request.POST.get('lieu_trouve', '').strip() or None,
            visiteur_id=request.POST.get('visiteur') or None,
            visite_id=request.POST.get('visite') or None,
            departement_id=request.POST.get('departement') or None,
            photo_objet=request.POST.get('photo_objet', '').strip() or None,
            created_by=request.user,
        )
        HistoriqueAction.objects.create(utilisateur=request.user, action='AJOUT', entite='ObjetOublie', entite_id=obj.pk)
        nom_visiteur = ''
        if obj.visite and obj.visite.visiteur:
            nom_visiteur = f'{obj.visite.visiteur.prenom} {obj.visite.visiteur.nom}'
        Alerte.objects.create(
            type='OBJET_OUBLIE',
            entite_id=obj.pk,
            message=f"Objet oublié — {nom_visiteur} — {obj.nom_objet}" if nom_visiteur else f"Objet oublié — {obj.nom_objet}",
        )
        messages.success(request, 'Objet oublié créé avec succès.')
        return redirect('liste_objets_oublies')
    return render(request, 'objets_oublies/ajouter.html', _base_ctx())


@login_required
def detail_objet_oublie(request, pk):
    item = get_object_or_404(
        ObjetOublie.objects.select_related(
            'visite__visiteur', 'departement', 'created_by', 'updated_by'
        ),
        pk=pk
    )
    return render(request, 'objets_oublies/detail.html', {'item': item})


@login_required
def modifier_objet_oublie(request, pk):
    item = get_object_or_404(ObjetOublie, pk=pk)
    if request.method == 'POST':
        item.nom_objet = request.POST.get('nom_objet', '').strip()
        item.categorie = request.POST.get('categorie', '').strip() or None
        item.description = request.POST.get('description', '').strip() or None
        item.date_trouve = request.POST.get('date_trouve', '').strip() or None
        item.lieu_trouve = request.POST.get('lieu_trouve', '').strip() or None
        item.visiteur_id = request.POST.get('visiteur') or None
        item.visite_id = request.POST.get('visite') or None
        item.departement_id = request.POST.get('departement') or None
        item.date_remise = request.POST.get('date_remise', '').strip() or None
        item.personne_recupere = request.POST.get('personne_recupere', '').strip() or None
        item.signature = request.POST.get('signature', '').strip() or None
        item.photo_objet = request.POST.get('photo_objet', '').strip() or None
        item.statut = request.POST.get('statut', 'NON_RESTITUÉ')
        item.updated_by = request.user
        item.save()
        HistoriqueAction.objects.create(utilisateur=request.user, action='MODIFICATION', entite='ObjetOublie', entite_id=item.pk)
        messages.success(request, 'Objet oublié modifié avec succès.')
        return redirect('liste_objets_oublies')
    return render(request, 'objets_oublies/modifier.html', {'item': item, **_base_ctx()})


@login_required
def supprimer_objet_oublie(request, pk):
    item = get_object_or_404(ObjetOublie, pk=pk)
    if request.method == 'POST':
        item.delete()
        HistoriqueAction.objects.create(utilisateur=request.user, action='SUPPRESSION', entite='ObjetOublie', entite_id=item.pk)
        messages.success(request, 'Objet oublié supprimé avec succès.')
        return redirect('liste_objets_oublies')
    return render(request, 'objets_oublies/detail.html', {'item': item, 'confirm_delete': True})


@login_required
def restituer_objet_oublie(request, pk):
    item = get_object_or_404(ObjetOublie, pk=pk)
    if request.method == 'POST':
        if item.statut == 'RESTITUÉ':
            messages.warning(request, 'Cet objet a déjà été restitué.')
            return redirect('liste_objets_oublies')
        item.personne_recupere = request.POST.get('personne_recupere', '').strip() or None
        item.signature = request.POST.get('signature', '').strip() or None
        item.date_remise = timezone.now()
        item.statut = 'RESTITUÉ'
        item.updated_by = request.user
        item.save(update_fields=['personne_recupere', 'signature', 'date_remise', 'statut', 'updated_by', 'updated_at'])
        HistoriqueAction.log(request, 'RESTITUTION', 'ObjetOublie', entite_id=item.pk,
                             details='Objet "{}" restitué à {}'.format(item.nom_objet, item.personne_recupere or '?'))
        messages.success(request, 'Objet "{}" marqué comme restitué.'.format(item.nom_objet))
    return redirect('liste_objets_oublies')


@login_required
def detection_objets_oublies(request):
    items = ObjetOublie.objects.select_related(
        'visite__visiteur', 'departement'
    ).order_by('-created_at')
    paginator = Paginator(items, 10)
    page_obj = paginator.get_page(request.GET.get('page', 1))
    page_links = paginator.get_elided_page_range(page_obj.number, on_each_side=2, on_ends=1)
    return render(request, 'objets_oublies/detection.html', {
        'page_obj': page_obj, 'page_links': page_links,
    })


@login_required
def export_pdf_objets_oublies(request):
    items = ObjetOublie.objects.select_related(
        'visite__visiteur', 'departement'
    ).order_by('-created_at')
    headers = ['N°', 'Objet', 'Catégorie', 'Visiteur', 'Département', 'Statut', 'Date trouvé']
    rows = [[i + 1, obj.nom_objet, obj.categorie or '-',
             str(obj.visite.visiteur) if obj.visite and obj.visite.visiteur else '-',
             obj.departement.nom if obj.departement else '-',
             obj.statut,
             obj.date_trouve.strftime('%d/%m/%Y %H:%M') if obj.date_trouve else '-']
            for i, obj in enumerate(items)]
    return _export_pdf(rows, headers, 'Liste des objets oubliés', 'objets_oublies', [0.5, 2, 1.5, 2, 1.5, 1, 1.5])


@login_required
def export_excel_objets_oublies(request):
    items = ObjetOublie.objects.select_related(
        'visite__visiteur', 'departement'
    ).order_by('-created_at')
    headers = ['N°', 'Objet', 'Catégorie', 'Visiteur', 'Département', 'Statut', 'Date trouvé']
    rows = [[i + 1, obj.nom_objet, obj.categorie or '',
             str(obj.visite.visiteur) if obj.visite and obj.visite.visiteur else '',
             obj.departement.nom if obj.departement else '',
             obj.statut,
             obj.date_trouve.strftime('%d/%m/%Y %H:%M') if obj.date_trouve else '']
            for i, obj in enumerate(items)]
    return _export_excel(rows, headers, 'Objets oubliés', 'objets_oublies')
