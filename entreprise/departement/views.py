from io import BytesIO

from django.http import HttpResponse
from django.shortcuts import render, redirect, get_object_or_404
from django.contrib.auth.decorators import login_required
from django.contrib import messages
from django.core.paginator import Paginator
from django.db.models import Q
from django.utils import timezone
from entreprise.models import Departement
from utilisateurs.models import HistoriqueAction
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.units import cm
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle


@login_required
def liste_departements(request):
    query = request.GET.get('q', '').strip()
    status_filter = request.GET.get('statut', 'TOUS')

    items = Departement.objects.all()

    if query:
        items = items.filter(
            Q(nom__icontains=query) |
            Q(description__icontains=query)
        )
    if status_filter != 'TOUS':
        items = items.filter(statut=status_filter)

    items = items.order_by('-created_at')

    paginator = Paginator(items, 10)
    page_number = request.GET.get('page', 1)
    page_obj = paginator.get_page(page_number)
    page_links = paginator.get_elided_page_range(page_obj.number, on_each_side=2, on_ends=1)

    context = {
        'page_obj': page_obj,
        'page_links': page_links,
        'query': query,
        'status_filter': status_filter,
    }
    return render(request, 'departements/liste.html', context)


@login_required
def ajouter_departement(request):
    if request.method == 'POST':
        nom = request.POST.get('nom', '').strip()
        if not nom:
            messages.error(request, 'Le nom du département est requis.')
            return render(request, 'departements/ajouter.html')

        if Departement.objects.filter(nom__iexact=nom).exists():
            messages.error(request, 'Un département avec ce nom existe déjà.')
            return render(request, 'departements/ajouter.html')

        d = Departement.objects.create(
            nom=nom,
            description=request.POST.get('description', '').strip() or None,
            created_by=request.user,
        )
        HistoriqueAction.log(request, 'AJOUT', 'Departement', entite_id=d.pk, details=nom)
        messages.success(request, 'Ajout effectué avec succès.')
        return redirect('liste_departements')

    return render(request, 'departements/ajouter.html')


@login_required
def detail_departement(request, pk):
    item = get_object_or_404(Departement.objects.select_related('created_by', 'updated_by'), pk=pk)
    return render(request, 'departements/detail.html', {'item': item})


@login_required
def modifier_departement(request, pk):
    item = get_object_or_404(Departement, pk=pk)

    if request.method == 'POST':
        nom = request.POST.get('nom', '').strip()
        if not nom:
            messages.error(request, 'Le nom du département est requis.')
            return render(request, 'departements/modifier.html', {'item': item})

        if Departement.objects.filter(nom__iexact=nom).exclude(pk=item.pk).exists():
            messages.error(request, 'Un département avec ce nom existe déjà.')
            return render(request, 'departements/modifier.html', {'item': item})

        item.nom = nom
        item.description = request.POST.get('description', '').strip() or None
        item.statut = request.POST.get('statut', 'ACTIF')
        item.updated_by = request.user
        item.save()
        HistoriqueAction.log(request, 'MODIFICATION', 'Departement', entite_id=item.pk, details=item.nom)
        messages.success(request, 'Modification effectuée avec succès.')
        return redirect('liste_departements')

    return render(request, 'departements/modifier.html', {'item': item})


@login_required
def export_pdf_departements(request):
    items = Departement.objects.all().order_by('-created_at')

    buffer = BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=landscape(A4), title="Départements",
                            leftMargin=1.5*cm, rightMargin=1.5*cm)
    styles = getSampleStyleSheet()
    cell_style = ParagraphStyle('CellStyle', parent=styles['Normal'],
                                fontSize=9, leading=12, alignment=1)
    elements = []

    elements.append(Paragraph("Départements", styles['Title']))
    elements.append(Spacer(1, 0.5 * cm))
    elements.append(Paragraph(f"Généré le {timezone.now().strftime('%d/%m/%Y %H:%M')}", styles['Normal']))
    elements.append(Spacer(1, 0.5 * cm))

    headers = ['N°', 'Département', 'Description', 'Statut', 'Créé le']
    col_ratios = [0.5, 2, 4, 1, 1.5]
    available_width = landscape(A4)[0] - 3 * cm
    total_ratio = sum(col_ratios)
    col_widths = [available_width * r / total_ratio for r in col_ratios]

    data = [headers]
    for i, p in enumerate(items, 1):
        data.append([
            Paragraph(str(i), cell_style),
            Paragraph(p.nom, cell_style),
            Paragraph(p.description or '-', cell_style),
            Paragraph(p.statut, cell_style),
            Paragraph(p.created_at.strftime('%d/%m/%Y %H:%M') if p.created_at else '-', cell_style),
        ])

    table = Table(data, colWidths=col_widths, repeatRows=1)
    table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1270b8')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, -1), 9),
        ('BOTTOMPADDING', (0, 0), (-1, 0), 10),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#e2e8f0')),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.HexColor('#f8fafc'), colors.white]),
        ('VALIGN', (0, 0), (-1, -1), 'TOP'),
    ]))
    elements.append(table)

    doc.build(elements)
    pdf = buffer.getvalue()
    buffer.close()

    response = HttpResponse(pdf, content_type='application/pdf')
    response['Content-Disposition'] = f'inline; filename="departements_{timezone.now().strftime("%Y%m%d")}.pdf"'
    return response


@login_required
def export_excel_departements(request):
    items = Departement.objects.all().order_by('-created_at')

    wb = Workbook()
    ws = wb.active
    ws.title = "Départements"
    ws.append(['N°', 'Département', 'Description', 'Statut', 'Créé le'])

    for i, p in enumerate(items, 1):
        ws.append([i, p.nom, p.description or '', p.statut, p.created_at.strftime('%d/%m/%Y %H:%M') if p.created_at else ''])

    ws.column_dimensions['A'].width = 6
    ws.column_dimensions['B'].width = 30
    ws.column_dimensions['C'].width = 40
    ws.column_dimensions['D'].width = 12
    ws.column_dimensions['E'].width = 18

    header_fill = PatternFill(start_color='1270b8', end_color='1270b8', fill_type='solid')
    header_font = Font(bold=True, color='ffffff', size=11)
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    response = HttpResponse(buffer, content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = f'attachment; filename="departements_{timezone.now().strftime("%Y%m%d")}.xlsx"'
    return response


@login_required
def supprimer_departement(request, pk):
    item = get_object_or_404(Departement, pk=pk)

    if request.method == 'POST':
        item.delete()
        HistoriqueAction.log(request, 'SUPPRESSION', 'Departement', entite_id=item.pk, details=item.nom)
        messages.success(request, 'Suppression effectuée avec succès.')
        return redirect('liste_departements')

    return render(request, 'departements/detail.html', {'item': item, 'confirm_delete': True})
