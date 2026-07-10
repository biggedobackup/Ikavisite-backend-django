from io import BytesIO

from django.http import HttpResponse
from django.shortcuts import render, redirect, get_object_or_404
from django.contrib.auth.decorators import login_required
from django.contrib import messages
from django.core.paginator import Paginator
from django.db.models import Q
from django.utils import timezone
from entreprise.models import PorteEntree
from utilisateurs.models import HistoriqueAction
from openpyxl import Workbook
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.units import cm
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle


@login_required
def liste_portes_entree(request):
    query = request.GET.get('q', '').strip()
    status_filter = request.GET.get('statut', 'TOUS')

    portes = PorteEntree.objects.all()

    if query:
        portes = portes.filter(
            Q(titre__icontains=query) |
            Q(emplacement__icontains=query) |
            Q(description__icontains=query)
        )
    if status_filter != 'TOUS':
        portes = portes.filter(statut=status_filter)

    portes = portes.order_by('-created_at')

    paginator = Paginator(portes, 10)
    page_number = request.GET.get('page', 1)
    page_obj = paginator.get_page(page_number)
    page_links = paginator.get_elided_page_range(page_obj.number, on_each_side=2, on_ends=1)

    context = {
        'page_obj': page_obj,
        'page_links': page_links,
        'query': query,
        'status_filter': status_filter,
    }
    return render(request, 'portes-entree/liste.html', context)


@login_required
def ajouter_porte_entree(request):
    if request.method == 'POST':
        titre = request.POST.get('titre', '').strip()
        if not titre:
            messages.error(request, 'Le nom de la porte est requis.')
            return render(request, 'portes-entree/ajouter.html')

        if PorteEntree.objects.filter(titre__iexact=titre).exists():
            messages.error(request, 'Une porte avec ce nom existe déjà.')
            return render(request, 'portes-entree/ajouter.html')

        p = PorteEntree.objects.create(
            titre=titre,
            emplacement=request.POST.get('emplacement', '').strip() or None,
            description=request.POST.get('description', '').strip() or None,
            created_by=request.user,
        )
        HistoriqueAction.log(request, 'AJOUT', 'PorteEntree', entite_id=p.pk, details=titre)
        messages.success(request, 'Ajout effectué avec succès.')
        return redirect('liste_portes_entree')

    return render(request, 'portes-entree/ajouter.html')


@login_required
def detail_porte_entree(request, pk):
    porte = get_object_or_404(PorteEntree.objects.select_related('created_by', 'updated_by'), pk=pk)
    return render(request, 'portes-entree/detail.html', {'porte': porte})


@login_required
def modifier_porte_entree(request, pk):
    porte = get_object_or_404(PorteEntree, pk=pk)

    if request.method == 'POST':
        titre = request.POST.get('titre', '').strip()
        if not titre:
            messages.error(request, 'Le nom de la porte est requis.')
            return render(request, 'portes-entree/modifier.html', {'porte': porte})

        if PorteEntree.objects.filter(titre__iexact=titre).exclude(pk=porte.pk).exists():
            messages.error(request, 'Une porte avec ce nom existe déjà.')
            return render(request, 'portes-entree/modifier.html', {'porte': porte})

        porte.titre = titre
        porte.emplacement = request.POST.get('emplacement', '').strip() or None
        porte.description = request.POST.get('description', '').strip() or None
        porte.statut = request.POST.get('statut', 'ACTIF')
        porte.updated_by = request.user
        porte.save()
        HistoriqueAction.log(request, 'MODIFICATION', 'PorteEntree', entite_id=porte.pk, details=porte.titre)
        messages.success(request, 'Modification effectuée avec succès.')
        return redirect('liste_portes_entree')

    return render(request, 'portes-entree/modifier.html', {'porte': porte})


@login_required
def export_pdf_portes_entree(request):
    portes = PorteEntree.objects.all().order_by('-created_at')

    buffer = BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=landscape(A4), title="Portes d'entrée",
                            leftMargin=1.5*cm, rightMargin=1.5*cm)
    styles = getSampleStyleSheet()
    cell_style = ParagraphStyle('CellStyle', parent=styles['Normal'],
                                fontSize=9, leading=12, alignment=1)
    elements = []

    elements.append(Paragraph("Portes d'entrée", styles['Title']))
    elements.append(Spacer(1, 0.5 * cm))
    elements.append(Paragraph(f"Généré le {timezone.now().strftime('%d/%m/%Y %H:%M')}", styles['Normal']))
    elements.append(Spacer(1, 0.5 * cm))

    headers = ['N°', 'Porte', 'Emplacement', 'Statut', 'Créée le']
    col_ratios = [0.5, 2, 2, 1, 1.5]
    available_width = landscape(A4)[0] - 3 * cm
    total_ratio = sum(col_ratios)
    col_widths = [available_width * r / total_ratio for r in col_ratios]

    data = [headers]
    for i, p in enumerate(portes, 1):
        data.append([
            Paragraph(str(i), cell_style),
            Paragraph(p.titre, cell_style),
            Paragraph(p.emplacement or '-', cell_style),
            Paragraph(p.statut, cell_style),
            Paragraph(p.created_at.strftime('%d/%m/%Y %H:%M') if p.created_at else '-', cell_style),
        ])

    table = Table(data, colWidths=col_widths, repeatRows=1)
    table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1270b8')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, -1), 9),
        ('BOTTOMPADDING', (0, 0), (-1, 0), 10),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#e2e8f0')),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.HexColor('#f8fafc'), colors.white]),
        ('VALIGN', (0, 0), (-1, -1), 'TOP'),
    ]))
    elements.append(table)

    doc.build(elements)
    pdf = buffer.getvalue()
    buffer.close()

    response = HttpResponse(pdf, content_type='application/pdf')
    response['Content-Disposition'] = f'inline; filename="portes_entree_{timezone.now().strftime("%Y%m%d")}.pdf"'
    return response


@login_required
def export_excel_portes_entree(request):
    portes = PorteEntree.objects.all().order_by('-created_at')

    wb = Workbook()
    ws = wb.active
    ws.title = "Portes d'entrée"
    ws.append(['N°', 'Porte', 'Emplacement', 'Description', 'Statut', 'Créée le'])

    for i, p in enumerate(portes, 1):
        ws.append([i, p.titre, p.emplacement or '', p.description or '', p.statut, p.created_at.strftime('%d/%m/%Y %H:%M') if p.created_at else ''])

    ws.column_dimensions['A'].width = 6
    ws.column_dimensions['B'].width = 30
    ws.column_dimensions['C'].width = 25
    ws.column_dimensions['D'].width = 40
    ws.column_dimensions['E'].width = 12
    ws.column_dimensions['F'].width = 18

    from openpyxl.styles import Font, PatternFill
    header_fill = PatternFill(start_color='1270b8', end_color='1270b8', fill_type='solid')
    header_font = Font(bold=True, color='ffffff', size=11)
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    response = HttpResponse(buffer, content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = f'attachment; filename="portes_entree_{timezone.now().strftime("%Y%m%d")}.xlsx"'
    return response


@login_required
def supprimer_porte_entree(request, pk):
    porte = get_object_or_404(PorteEntree, pk=pk)

    if request.method == 'POST':
        porte.delete()
        HistoriqueAction.log(request, 'SUPPRESSION', 'PorteEntree', entite_id=porte.pk, details=porte.titre)
        messages.success(request, 'Suppression effectuée avec succès.')
        return redirect('liste_portes_entree')

    return render(request, 'portes-entree/detail.html', {'porte': porte, 'confirm_delete': True})
