import csv
from io import BytesIO, StringIO

from django.http import HttpResponse
from django.shortcuts import render, redirect, get_object_or_404
from django.contrib.auth.decorators import login_required
from django.contrib import messages
from django.core.paginator import Paginator
from django.db.models import Q
from django.utils import timezone
from entreprise.models import Personnel, Departement
from utilisateurs.models import HistoriqueAction
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.units import cm
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle


def _get_departements():
    return Departement.objects.filter(statut='ACTIF').order_by('nom')


@login_required
def liste_personnel(request):
    query = request.GET.get('q', '').strip()
    status_filter = request.GET.get('statut', 'TOUS')
    dept_filter = request.GET.get('departement', 'TOUS')

    items = Personnel.objects.all().select_related('departement')

    if query:
        items = items.filter(
            Q(nom__icontains=query) |
            Q(prenom__icontains=query) |
            Q(fonction__icontains=query) |
            Q(email__icontains=query) |
            Q(telephone__icontains=query) |
            Q(departement__nom__icontains=query)
        )
    if status_filter != 'TOUS':
        items = items.filter(statut=status_filter)
    if dept_filter != 'TOUS':
        items = items.filter(departement__nom=dept_filter)

    items = items.order_by('-created_at')

    paginator = Paginator(items, 10)
    page_number = request.GET.get('page', 1)
    page_obj = paginator.get_page(page_number)
    page_links = paginator.get_elided_page_range(page_obj.number, on_each_side=2, on_ends=1)

    from urllib.parse import urlencode
    fp = {k: v for k, v in [('q', query), ('departement', dept_filter), ('statut', status_filter)]
          if v and v != 'TOUS'}
    context = {
        'page_obj': page_obj,
        'page_links': page_links,
        'query': query,
        'status_filter': status_filter,
        'dept_filter': dept_filter,
        'departements': _get_departements(),
        'filter_params': urlencode(fp) + '&' if fp else '',
    }
    return render(request, 'personnel/liste.html', context)


@login_required
def ajouter_personnel(request):
    if request.method == 'POST':
        nom = request.POST.get('nom', '').strip()
        prenom = request.POST.get('prenom', '').strip()

        if not nom or not prenom:
            messages.error(request, 'Le nom et le prénom sont requis.')
            return render(request, 'personnel/ajouter.html', {'departements': _get_departements()})

        dept_id = request.POST.get('id_departement', '').strip()
        departement = None
        if dept_id:
            try:
                departement = Departement.objects.get(pk=int(dept_id))
            except (ValueError, Departement.DoesNotExist):
                pass

        p = Personnel.objects.create(
            nom=nom,
            prenom=prenom,
            fonction=request.POST.get('fonction', '').strip() or None,
            email=request.POST.get('email', '').strip() or None,
            telephone=request.POST.get('telephone', '').strip() or None,
            departement=departement,
            created_by=request.user,
        )
        HistoriqueAction.log(request, 'AJOUT', 'Personnel', entite_id=p.pk, details=f'{prenom} {nom}')
        messages.success(request, 'Ajout effectué avec succès.')
        return redirect('liste_personnel')

    return render(request, 'personnel/ajouter.html', {'departements': _get_departements()})


@login_required
def detail_personnel(request, pk):
    item = get_object_or_404(
        Personnel.objects.select_related('departement', 'created_by', 'updated_by'),
        pk=pk
    )
    return render(request, 'personnel/detail.html', {'item': item})


@login_required
def modifier_personnel(request, pk):
    item = get_object_or_404(Personnel, pk=pk)

    if request.method == 'POST':
        nom = request.POST.get('nom', '').strip()
        prenom = request.POST.get('prenom', '').strip()

        if not nom or not prenom:
            messages.error(request, 'Le nom et le prénom sont requis.')
            return render(request, 'personnel/modifier.html', {'item': item, 'departements': _get_departements()})

        dept_id = request.POST.get('id_departement', '').strip()
        departement = None
        if dept_id:
            try:
                departement = Departement.objects.get(pk=int(dept_id))
            except (ValueError, Departement.DoesNotExist):
                pass

        item.nom = nom
        item.prenom = prenom
        item.fonction = request.POST.get('fonction', '').strip() or None
        item.email = request.POST.get('email', '').strip() or None
        item.telephone = request.POST.get('telephone', '').strip() or None
        item.departement = departement
        item.statut = request.POST.get('statut', 'ACTIF')
        item.updated_by = request.user
        item.save()
        HistoriqueAction.log(request, 'MODIFICATION', 'Personnel', entite_id=item.pk, details=f'{prenom} {nom}')
        messages.success(request, 'Modification effectuée avec succès.')
        return redirect('liste_personnel')

    return render(request, 'personnel/modifier.html', {'item': item, 'departements': _get_departements()})


@login_required
def export_pdf_personnel(request):
    items = Personnel.objects.all().select_related('departement').order_by('-created_at')

    buffer = BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=landscape(A4), title="Personnel",
                            leftMargin=1.5*cm, rightMargin=1.5*cm)
    styles = getSampleStyleSheet()
    cell_style = ParagraphStyle('CellStyle', parent=styles['Normal'],
                                fontSize=8, leading=10, alignment=1)
    elements = []

    elements.append(Paragraph("Personnel", styles['Title']))
    elements.append(Spacer(1, 0.5 * cm))
    elements.append(Paragraph(f"Généré le {timezone.now().strftime('%d/%m/%Y %H:%M')}", styles['Normal']))
    elements.append(Spacer(1, 0.5 * cm))

    headers = ['N°', 'Prénom', 'Nom', 'Fonction', 'Email', 'Téléphone', 'Département', 'Statut']
    col_ratios = [0.5, 1.5, 1.5, 2, 2, 1.5, 1.5, 1]
    available_width = landscape(A4)[0] - 3 * cm
    total_ratio = sum(col_ratios)
    col_widths = [available_width * r / total_ratio for r in col_ratios]

    data = [headers]
    for i, p in enumerate(items, 1):
        data.append([
            Paragraph(str(i), cell_style),
            Paragraph(p.prenom, cell_style),
            Paragraph(p.nom, cell_style),
            Paragraph(p.fonction or '-', cell_style),
            Paragraph(p.email or '-', cell_style),
            Paragraph(p.telephone or '-', cell_style),
            Paragraph(p.departement.nom if p.departement else '-', cell_style),
            Paragraph(p.statut, cell_style),
        ])

    table = Table(data, colWidths=col_widths, repeatRows=1)
    table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1270b8')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, -1), 8),
        ('BOTTOMPADDING', (0, 0), (-1, 0), 10),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#e2e8f0')),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.HexColor('#f8fafc'), colors.white]),
        ('VALIGN', (0, 0), (-1, -1), 'TOP'),
    ]))
    elements.append(table)

    doc.build(elements)
    pdf = buffer.getvalue()
    buffer.close()

    response = HttpResponse(pdf, content_type='application/pdf')
    response['Content-Disposition'] = f'inline; filename="personnel_{timezone.now().strftime("%Y%m%d")}.pdf"'
    return response


@login_required
def export_excel_personnel(request):
    items = Personnel.objects.all().select_related('departement').order_by('-created_at')

    wb = Workbook()
    ws = wb.active
    ws.title = "Personnel"
    ws.append(['N°', 'Prénom', 'Nom', 'Fonction', 'Email', 'Téléphone', 'Département', 'Statut'])

    for i, p in enumerate(items, 1):
        ws.append([i, p.prenom, p.nom, p.fonction or '', p.email or '', p.telephone or '', p.departement.nom if p.departement else '', p.statut])

    ws.column_dimensions['A'].width = 6
    ws.column_dimensions['B'].width = 18
    ws.column_dimensions['C'].width = 18
    ws.column_dimensions['D'].width = 22
    ws.column_dimensions['E'].width = 28
    ws.column_dimensions['F'].width = 16
    ws.column_dimensions['G'].width = 22
    ws.column_dimensions['H'].width = 10

    header_fill = PatternFill(start_color='1270b8', end_color='1270b8', fill_type='solid')
    header_font = Font(bold=True, color='ffffff', size=11)
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    response = HttpResponse(buffer, content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = f'attachment; filename="personnel_{timezone.now().strftime("%Y%m%d")}.xlsx"'
    return response


@login_required
def telecharger_modele_csv_personnel(request):
    headers = ['prenom', 'nom', 'fonction', 'email', 'telephone', 'departement']
    sample = ['Jean', 'Dupont', 'Agent', 'jean.dupont@email.com', '0612345678', 'Informatique']
    buffer = StringIO()
    writer = csv.writer(buffer)
    writer.writerow(headers)
    writer.writerow(sample)
    response = HttpResponse(buffer.getvalue(), content_type='text/csv;charset=utf-8')
    response['Content-Disposition'] = 'attachment; filename="modele_personnel.csv"'
    return response


@login_required
def import_csv_personnel(request):
    if request.method != 'POST':
        messages.error(request, 'Méthode non autorisée.')
        return redirect('liste_personnel')
    file = request.FILES.get('file')
    if not file:
        messages.error(request, 'Aucun fichier fourni.')
        return redirect('liste_personnel')
    try:
        decoded = file.read().decode('utf-8')
        reader = csv.DictReader(StringIO(decoded))
        imported = 0
        errors = 0
        dept_cache = {d.nom: d for d in Departement.objects.all()}
        for row in reader:
            prenom = row.get('prenom', '').strip()
            nom = row.get('nom', '').strip()
            if not prenom or not nom:
                errors += 1
                continue
            departement = dept_cache.get(row.get('departement', '').strip())
            Personnel.objects.create(
                nom=nom,
                prenom=prenom,
                fonction=row.get('fonction', '').strip() or None,
                email=row.get('email', '').strip() or None,
                telephone=row.get('telephone', '').strip() or None,
                departement=departement,
                created_by=request.user,
            )
            imported += 1
        HistoriqueAction.log(request, 'IMPORT', 'Personnel', details=f'{imported} membres importés')
        messages.success(request, f'{imported} membre(s) importé(s) avec succès. {errors} erreur(s).')
    except Exception as e:
        messages.error(request, f'Erreur lors de l\'import: {str(e)}')
    return redirect('liste_personnel')


@login_required
def supprimer_personnel(request, pk):
    item = get_object_or_404(Personnel, pk=pk)

    if request.method == 'POST':
        item.delete()
        HistoriqueAction.log(request, 'SUPPRESSION', 'Personnel', entite_id=item.pk, details=f'{item.prenom} {item.nom}')
        messages.success(request, 'Suppression effectuée avec succès.')
        return redirect('liste_personnel')

    return render(request, 'personnel/detail.html', {'item': item, 'confirm_delete': True})
