from io import BytesIO

from django.http import HttpResponse
from django.shortcuts import render, redirect, get_object_or_404
from django.contrib.auth.decorators import login_required
from django.contrib import messages
from django.core.paginator import Paginator
from django.db.models import Q
from django.utils import timezone
from entreprise.models import CreneauSemaine
from utilisateurs.models import HistoriqueAction
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.units import cm
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle

JOURS = ['LUNDI', 'MARDI', 'MERCREDI', 'JEUDI', 'VENDREDI', 'SAMEDI', 'DIMANCHE']


@login_required
def liste_creneaux(request):
    query = request.GET.get('q', '').strip()

    # Grouper tous les créneaux par jour
    all_items = CreneauSemaine.objects.all().order_by('heure_debut')

    jours_data = []
    for jour in JOURS:
        slots = [s for s in all_items if s.jour_semaine == jour]
        # Un jour est "actif" s'il a au moins un créneau ACTIF
        has_active = any(s.statut == 'ACTIF' for s in slots)
        jours_data.append({
            'jour_semaine': jour,
            'is_active': has_active,
            'slots': slots,
        })

    context = {
        'jours_data': jours_data,
        'query': query,
    }
    return render(request, 'creneaux-semaine/liste.html', context)


@login_required
def ajouter_creneau(request):
    if request.method == 'POST':
        jour = request.POST.get('jour_semaine', '').strip()
        heure_debut = request.POST.get('heure_debut', '').strip()
        heure_fin = request.POST.get('heure_fin', '').strip()

        if not jour or not heure_debut or not heure_fin:
            messages.error(request, 'Tous les champs sont requis.')
            return render(request, 'creneaux-semaine/ajouter.html', {'jours': JOURS})

        c = CreneauSemaine.objects.create(
            jour_semaine=jour,
            heure_debut=heure_debut,
            heure_fin=heure_fin,
            created_by=request.user,
        )
        HistoriqueAction.log(request, 'AJOUT', 'CreneauSemaine', entite_id=c.pk, details=f'{jour} {heure_debut}-{heure_fin}')
        messages.success(request, 'Ajout effectué avec succès.')
        return redirect('liste_creneaux')

    return render(request, 'creneaux-semaine/ajouter.html', {'jours': JOURS})


@login_required
def detail_creneau(request, pk):
    item = get_object_or_404(CreneauSemaine.objects.select_related('created_by', 'updated_by'), pk=pk)
    return render(request, 'creneaux-semaine/detail.html', {'item': item})


@login_required
def modifier_creneau(request, pk):
    item = get_object_or_404(CreneauSemaine, pk=pk)

    if request.method == 'POST':
        jour = request.POST.get('jour_semaine', '').strip()
        heure_debut = request.POST.get('heure_debut', '').strip()
        heure_fin = request.POST.get('heure_fin', '').strip()

        if not jour or not heure_debut or not heure_fin:
            messages.error(request, 'Tous les champs sont requis.')
            return render(request, 'creneaux-semaine/modifier.html', {'item': item, 'jours': JOURS})

        item.jour_semaine = jour
        item.heure_debut = heure_debut
        item.heure_fin = heure_fin
        item.statut = request.POST.get('statut', 'ACTIF')
        item.updated_by = request.user
        item.save()
        HistoriqueAction.log(request, 'MODIFICATION', 'CreneauSemaine', entite_id=item.pk, details=f'{item.jour_semaine} {item.heure_debut}-{item.heure_fin}')
        messages.success(request, 'Modification effectuée avec succès.')
        return redirect('liste_creneaux')

    return render(request, 'creneaux-semaine/modifier.html', {'item': item, 'jours': JOURS})


@login_required
def export_pdf_creneaux(request):
    items = CreneauSemaine.objects.all().order_by('-created_at')

    buffer = BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=landscape(A4), title="Créneaux de semaine",
                            leftMargin=1.5*cm, rightMargin=1.5*cm)
    styles = getSampleStyleSheet()
    cell_style = ParagraphStyle('CellStyle', parent=styles['Normal'],
                                fontSize=9, leading=12, alignment=1)
    elements = []

    elements.append(Paragraph("Créneaux de semaine", styles['Title']))
    elements.append(Spacer(1, 0.5 * cm))
    elements.append(Paragraph(f"Généré le {timezone.now().strftime('%d/%m/%Y %H:%M')}", styles['Normal']))
    elements.append(Spacer(1, 0.5 * cm))

    headers = ['N°', 'Jour', 'Début', 'Fin', 'Statut', 'Créé le']
    col_ratios = [0.5, 1.5, 1, 1, 1, 1.5]
    available_width = landscape(A4)[0] - 3 * cm
    total_ratio = sum(col_ratios)
    col_widths = [available_width * r / total_ratio for r in col_ratios]

    data = [headers]
    for i, p in enumerate(items, 1):
        data.append([
            Paragraph(str(i), cell_style),
            Paragraph(p.jour_semaine, cell_style),
            Paragraph(str(p.heure_debut), cell_style),
            Paragraph(str(p.heure_fin), cell_style),
            Paragraph(p.statut, cell_style),
            Paragraph(p.created_at.strftime('%d/%m/%Y %H:%M') if p.created_at else '-', cell_style),
        ])

    table = Table(data, colWidths=col_widths, repeatRows=1)
    table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1270b8')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, -1), 9),
        ('BOTTOMPADDING', (0, 0), (-1, 0), 10),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#e2e8f0')),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.HexColor('#f8fafc'), colors.white]),
        ('VALIGN', (0, 0), (-1, -1), 'TOP'),
    ]))
    elements.append(table)

    doc.build(elements)
    pdf = buffer.getvalue()
    buffer.close()

    response = HttpResponse(pdf, content_type='application/pdf')
    response['Content-Disposition'] = f'inline; filename="creneaux_{timezone.now().strftime("%Y%m%d")}.pdf"'
    return response


@login_required
def export_excel_creneaux(request):
    items = CreneauSemaine.objects.all().order_by('-created_at')

    wb = Workbook()
    ws = wb.active
    ws.title = "Créneaux"
    ws.append(['N°', 'Jour', 'Début', 'Fin', 'Statut', 'Créé le'])

    for i, p in enumerate(items, 1):
        ws.append([i, p.jour_semaine, str(p.heure_debut), str(p.heure_fin), p.statut, p.created_at.strftime('%d/%m/%Y %H:%M') if p.created_at else ''])

    ws.column_dimensions['A'].width = 6
    ws.column_dimensions['B'].width = 20
    ws.column_dimensions['C'].width = 12
    ws.column_dimensions['D'].width = 12
    ws.column_dimensions['E'].width = 12
    ws.column_dimensions['F'].width = 18

    header_fill = PatternFill(start_color='1270b8', end_color='1270b8', fill_type='solid')
    header_font = Font(bold=True, color='ffffff', size=11)
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    response = HttpResponse(buffer, content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = f'attachment; filename="creneaux_{timezone.now().strftime("%Y%m%d")}.xlsx"'
    return response


@login_required
def supprimer_creneau(request, pk):
    item = get_object_or_404(CreneauSemaine, pk=pk)

    if request.method == 'POST':
        item.delete()
        HistoriqueAction.log(request, 'SUPPRESSION', 'CreneauSemaine', entite_id=item.pk, details=f'{item.jour_semaine} {item.heure_debut}-{item.heure_fin}')
        messages.success(request, 'Suppression effectuée avec succès.')
        return redirect('liste_creneaux')

    return render(request, 'creneaux-semaine/detail.html', {'item': item, 'confirm_delete': True})


# ─── AJAX ────────────────────────────────────────────────────────────────────────


@login_required
def ajouter_creneau_ajax(request):
    """Ajout rapide d'un créneau, retour JSON."""
    from django.http import JsonResponse
    if request.method != 'POST':
        return JsonResponse({'ok': False, 'error': 'POST requis'}, status=405)

    jour = request.POST.get('jour_semaine', '').strip()
    heure_debut = request.POST.get('heure_debut', '').strip()
    heure_fin = request.POST.get('heure_fin', '').strip()

    if not jour or not heure_debut or not heure_fin:
        return JsonResponse({'ok': False, 'error': 'Champs requis'}, status=400)

    if heure_debut >= heure_fin:
        return JsonResponse({'ok': False, 'error': 'Début avant fin requis'}, status=400)

    c = CreneauSemaine.objects.create(
        jour_semaine=jour.upper(),
        heure_debut=heure_debut,
        heure_fin=heure_fin,
        created_by=request.user,
    )
    HistoriqueAction.log(request, 'AJOUT', 'CreneauSemaine',
                         entite_id=c.pk, details=f'{jour} {heure_debut}-{heure_fin}')
    return JsonResponse({'ok': True, 'id': c.pk})


@login_required
def supprimer_creneau_ajax(request):
    """Suppression rapide, retour JSON."""
    from django.http import JsonResponse
    if request.method != 'POST':
        return JsonResponse({'ok': False, 'error': 'POST requis'}, status=405)

    try:
        c = CreneauSemaine.objects.get(pk=request.POST.get('pk'))
    except CreneauSemaine.DoesNotExist:
        return JsonResponse({'ok': False, 'error': 'Créneau introuvable'}, status=404)

    HistoriqueAction.log(request, 'SUPPRESSION', 'CreneauSemaine',
                         entite_id=c.pk, details=f'{c.jour_semaine} {c.heure_debut}-{c.heure_fin}')
    c.delete()
    return JsonResponse({'ok': True})


@login_required
def sauvegarder_config_planning(request):
    """Sauvegarde de la configuration planning complète (batch)."""
    import json
    from django.http import JsonResponse

    if request.method != 'POST':
        return JsonResponse({'ok': False, 'error': 'POST requis'}, status=405)

    try:
        body = json.loads(request.body)
    except (ValueError, AttributeError):
        return JsonResponse({'ok': False, 'error': 'JSON invalide'}, status=400)

    days = body.get('days', [])
    for d in days:
        jour = d.get('jour_semaine', '').upper()
        is_active = d.get('is_active', False)

        # Mettre à jour les créneaux de ce jour
        CreneauSemaine.objects.filter(jour_semaine=jour).update(
            statut='ACTIF' if is_active else 'INACTIF'
        )

    HistoriqueAction.log(request, 'CONFIG_PLANNING', 'CreneauSemaine',
                         details=f'Sauvegarde planning: {len(days)} jours')

    return JsonResponse({'ok': True})
