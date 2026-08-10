from io import BytesIO
from urllib.parse import urlencode

from django.contrib.auth.decorators import login_required
from django.contrib import messages
from django.core.paginator import Paginator
from django.db.models import Q
from django.http import HttpResponse
from django.shortcuts import render, redirect, get_object_or_404
from django.utils import timezone

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.units import cm
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle

from .models import TypeListeNoire, ListeNoire, DetectionListeNoire
from entreprise.models import PorteEntree
from utilisateurs.models import HistoriqueAction


def _export_pdf(rows, headers, title, filename, col_ratios=None):
    buffer = BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=landscape(A4), title=title,
                            leftMargin=1.5*cm, rightMargin=1.5*cm)
    styles = getSampleStyleSheet()
    cell_style = ParagraphStyle('CellStyle', parent=styles['Normal'],
                                fontSize=8, leading=10, alignment=1)
    elements = []
    elements.append(Paragraph(title, styles['Title']))
    elements.append(Spacer(1, 0.5 * cm))
    elements.append(Paragraph(f"Généré le {timezone.now().strftime('%d/%m/%Y %H:%M')}", styles['Normal']))
    elements.append(Spacer(1, 0.5 * cm))
    col_count = len(headers)
    available_width = landscape(A4)[0] - 3 * cm
    if col_ratios is None:
        col_ratios = [1] * col_count
    total_ratio = sum(col_ratios)
    col_widths = [available_width * r / total_ratio for r in col_ratios]
    wrapped_rows = [[Paragraph(str(c), cell_style) for c in row] for row in rows]
    data = [headers] + wrapped_rows
    table = Table(data, colWidths=col_widths, repeatRows=1)
    table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1270b8')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, -1), 8),
        ('BOTTOMPADDING', (0, 0), (-1, 0), 8),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#e2e8f0')),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.HexColor('#f8fafc'), colors.white]),
        ('VALIGN', (0, 0), (-1, -1), 'TOP'),
    ]))
    elements.append(table)
    doc.build(elements)
    pdf = buffer.getvalue()
    buffer.close()
    response = HttpResponse(pdf, content_type='application/pdf')
    response['Content-Disposition'] = f'inline; filename="{filename}_{timezone.now().strftime("%Y%m%d")}.pdf"'
    return response


def _export_excel(rows, headers, title, filename):
    wb = Workbook()
    ws = wb.active
    ws.title = title[:31]
    ws.append(headers)
    for row in rows:
        ws.append(row)
    for col_idx in range(1, len(headers) + 1):
        ws.column_dimensions[chr(64 + col_idx)].width = 22
    header_fill = PatternFill(start_color='1270b8', end_color='1270b8', fill_type='solid')
    header_font = Font(bold=True, color='ffffff', size=11)
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    response = HttpResponse(buffer, content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = f'attachment; filename="{filename}_{timezone.now().strftime("%Y%m%d")}.xlsx"'
    return response


# ─── Types de liste noire ─────────────────────────────────────────

@login_required
def liste_types_liste_noire(request):
    query = request.GET.get('q', '').strip()
    statut = request.GET.get('statut', '')
    items = TypeListeNoire.objects.all()
    if query:
        items = items.filter(Q(nom__icontains=query) | Q(description__icontains=query))
    if statut and statut != 'TOUS':
        items = items.filter(statut=statut)
    items = items.order_by('nom')
    paginator = Paginator(items, 10)
    page_obj = paginator.get_page(request.GET.get('page', 1))
    page_links = paginator.get_elided_page_range(page_obj.number, on_each_side=2, on_ends=1)
    fp = {k: v for k, v in [('q', query), ('statut', statut)] if v and v != 'TOUS'}
    return render(request, 'liste_noire/types/liste.html', {
        'page_obj': page_obj, 'page_links': page_links, 'query': query, 'statut_filter': statut,
        'filter_params': urlencode(fp) + '&' if fp else '',
    })


@login_required
def ajouter_type_liste_noire(request):
    if request.method == 'POST':
        nom = request.POST.get('nom', '').strip()
        if not nom:
            messages.error(request, 'Le nom est requis.')
            return render(request, 'liste_noire/types/ajouter.html')
        obj = TypeListeNoire.objects.create(
            nom=nom,
            description=request.POST.get('description', '').strip() or None,
            niveau_risque=request.POST.get('niveau_risque', 'FAIBLE'),
            statut='ACTIF',
            created_by=request.user,
        )
        HistoriqueAction.log(request, action='AJOUT', entite='TypeListeNoire', entite_id=obj.pk, details=f'Type créé : {nom}')
        messages.success(request, 'Type de liste noire créé avec succès.')
        return redirect('liste_types_liste_noire')
    return render(request, 'liste_noire/types/ajouter.html')


@login_required
def detail_type_liste_noire(request, pk):
    item = get_object_or_404(TypeListeNoire, pk=pk)
    return render(request, 'liste_noire/types/detail.html', {'item': item})


@login_required
def modifier_type_liste_noire(request, pk):
    item = get_object_or_404(TypeListeNoire, pk=pk)
    if request.method == 'POST':
        item.nom = request.POST.get('nom', '').strip()
        item.description = request.POST.get('description', '').strip() or None
        item.niveau_risque = request.POST.get('niveau_risque', 'FAIBLE')
        item.statut = request.POST.get('statut', 'ACTIF')
        item.updated_by = request.user
        item.save()
        HistoriqueAction.log(request, action='MODIFICATION', entite='TypeListeNoire', entite_id=item.pk, details=f'Type modifié : {item.nom}')
        messages.success(request, 'Type de liste noire modifié avec succès.')
        return redirect('liste_types_liste_noire')
    return render(request, 'liste_noire/types/modifier.html', {'item': item})


@login_required
def supprimer_type_liste_noire(request, pk):
    item = get_object_or_404(TypeListeNoire, pk=pk)
    if request.method == 'POST':
        item.delete()
        HistoriqueAction.log(request, action='SUPPRESSION', entite='TypeListeNoire', entite_id=item.pk, details=f'Type supprimé : {item.nom}')
        messages.success(request, 'Type de liste noire supprimé.')
    return redirect('liste_types_liste_noire')


@login_required
def export_pdf_types_liste_noire(request):
    items = TypeListeNoire.objects.all().order_by('nom')
    headers = ['N°', 'Nom', 'Description', 'Niveau risque', 'Statut', 'Créé le']
    rows = [[i + 1, t.nom, t.description or '-', t.niveau_risque, t.statut,
             t.created_at.strftime('%d/%m/%Y') if t.created_at else '-']
            for i, t in enumerate(items)]
    return _export_pdf(rows, headers, 'Types de liste noire', 'types_liste_noire', [0.5, 2, 3, 1, 1, 1.5])


@login_required
def export_excel_types_liste_noire(request):
    items = TypeListeNoire.objects.all().order_by('nom')
    headers = ['N°', 'Nom', 'Description', 'Niveau risque', 'Statut', 'Créé le']
    rows = [[i + 1, t.nom, t.description or '', t.niveau_risque, t.statut,
             t.created_at.strftime('%d/%m/%Y') if t.created_at else '']
            for i, t in enumerate(items)]
    return _export_excel(rows, headers, 'Types liste noire', 'types_liste_noire')


# ─── Listes noires ────────────────────────────────────────────────

@login_required
def liste_listes_noires(request):
    query = request.GET.get('q', '').strip()
    statut = request.GET.get('statut', '')
    items = ListeNoire.objects.select_related('type_liste_noire')
    if query:
        items = items.filter(
            Q(nom__icontains=query) | Q(prenom__icontains=query) |
            Q(motif__icontains=query) | Q(numero_piece__icontains=query) |
            Q(type_liste_noire__nom__icontains=query)
        )
    if statut and statut != 'TOUS':
        items = items.filter(statut=statut)
    items = items.order_by('-created_at')
    paginator = Paginator(items, 10)
    page_obj = paginator.get_page(request.GET.get('page', 1))
    page_links = paginator.get_elided_page_range(page_obj.number, on_each_side=2, on_ends=1)
    fp = {k: v for k, v in [('q', query), ('statut', statut)] if v and v != 'TOUS'}
    return render(request, 'liste_noire/listes/liste.html', {
        'page_obj': page_obj, 'page_links': page_links, 'query': query, 'statut_filter': statut,
        'filter_params': urlencode(fp) + '&' if fp else '',
    })


@login_required
def ajouter_liste_noire(request):
    type_defaut = TypeListeNoire.objects.filter(statut='ACTIF').first()
    type_liste_noire_id = type_defaut.pk if type_defaut else None
    if request.method == 'POST':
        nom = request.POST.get('nom', '').strip()
        prenom = request.POST.get('prenom', '').strip()
        if not nom or not prenom:
            messages.error(request, 'Le nom et le prénom sont requis.')
            return render(request, 'liste_noire/listes/ajouter.html')
        obj = ListeNoire.objects.create(
            type_liste_noire_id=type_liste_noire_id,
            nom=nom, prenom=prenom,
            motif=request.POST.get('motif', '').strip() or None,
            piece_identite=request.POST.get('piece_identite', '').strip() or None,
            numero_piece=request.POST.get('numero_piece', '').strip() or None,
            numero_nip=request.POST.get('numero_nip', '').strip() or None,
            date_debut=request.POST.get('date_debut', '').strip() or None,
            date_fin=request.POST.get('date_fin', '').strip() or None,
            blocage_automatique=request.POST.get('blocage_automatique') == 'on',
            created_by=request.user,
        )
        HistoriqueAction.log(request, action='AJOUT', entite='ListeNoire', entite_id=obj.pk, details=f'Inscrit : {nom} {prenom}')
        messages.success(request, 'Inscription ajoutée avec succès.')
        return redirect('liste_listes_noires')
    return render(request, 'liste_noire/listes/ajouter.html')


@login_required
def detail_liste_noire(request, pk):
    item = get_object_or_404(
        ListeNoire.objects.select_related(
            'type_liste_noire', 'created_by', 'updated_by'
        ),
        pk=pk
    )
    return render(request, 'liste_noire/listes/detail.html', {'item': item})


@login_required
def modifier_liste_noire(request, pk):
    item = get_object_or_404(ListeNoire, pk=pk)
    if request.method == 'POST':
        item.nom = request.POST.get('nom', '').strip()
        item.prenom = request.POST.get('prenom', '').strip()
        item.motif = request.POST.get('motif', '').strip() or None
        item.piece_identite = request.POST.get('piece_identite', '').strip() or None
        item.numero_piece = request.POST.get('numero_piece', '').strip() or None
        item.numero_nip = request.POST.get('numero_nip', '').strip() or None
        item.date_debut = request.POST.get('date_debut', '').strip() or None
        item.date_fin = request.POST.get('date_fin', '').strip() or None
        item.statut = request.POST.get('statut', 'ACTIF')
        item.blocage_automatique = request.POST.get('blocage_automatique') == 'on'
        item.updated_by = request.user
        item.save()
        HistoriqueAction.log(request, action='MODIFICATION', entite='ListeNoire', entite_id=item.pk, details=f'Inscription modifiée : {item.nom} {item.prenom}')
        messages.success(request, 'Inscription modifiée avec succès.')
        return redirect('liste_listes_noires')
    return render(request, 'liste_noire/listes/modifier.html', {'item': item})


@login_required
def supprimer_liste_noire(request, pk):
    item = get_object_or_404(ListeNoire, pk=pk)
    if request.method == 'POST':
        item.delete()
        HistoriqueAction.log(request, action='SUPPRESSION', entite='ListeNoire', entite_id=item.pk, details=f'Inscription supprimée : {item.nom} {item.prenom}')
        messages.success(request, 'Inscription supprimée.')
    return redirect('liste_listes_noires')


@login_required
def toggle_statut_liste_noire(request, pk):
    item = get_object_or_404(ListeNoire, pk=pk)
    if request.method == 'POST':
        item.statut = 'INACTIF' if item.statut == 'ACTIF' else 'ACTIF'
        item.updated_by = request.user
        item.save()
        HistoriqueAction.log(request, action='MODIFICATION', entite='ListeNoire', entite_id=item.pk, details=f'Statut changé à {item.statut} pour {item.nom} {item.prenom}')
        messages.success(request, f'Statut changé à {item.statut}.')
    return redirect('liste_listes_noires')


@login_required
def export_pdf_listes_noires(request):
    items = ListeNoire.objects.select_related('type_liste_noire').order_by('-created_at')
    headers = ['N°', 'Type', 'Nom', 'Prénom', 'N° Pièce', 'Motif', 'Début', 'Fin', 'Statut']
    rows = [[i + 1, ln.type_liste_noire.nom if ln.type_liste_noire else '-',
             ln.nom, ln.prenom, ln.numero_piece or '-', ln.motif or '-',
             ln.date_debut.strftime('%d/%m/%Y') if ln.date_debut else '-',
             ln.date_fin.strftime('%d/%m/%Y') if ln.date_fin else '-', ln.statut]
            for i, ln in enumerate(items)]
    return _export_pdf(rows, headers, 'Listes noires', 'listes_noires', [0.5, 1.5, 1.5, 1.5, 1.5, 2, 1, 1, 1])


@login_required
def export_excel_listes_noires(request):
    items = ListeNoire.objects.select_related('type_liste_noire').order_by('-created_at')
    headers = ['N°', 'Type', 'Nom', 'Prénom', 'N° Pièce', 'Motif', 'Début', 'Fin', 'Statut']
    rows = [[i + 1, ln.type_liste_noire.nom if ln.type_liste_noire else '',
             ln.nom, ln.prenom, ln.numero_piece or '', ln.motif or '',
             ln.date_debut.strftime('%d/%m/%Y') if ln.date_debut else '',
             ln.date_fin.strftime('%d/%m/%Y') if ln.date_fin else '', ln.statut]
            for i, ln in enumerate(items)]
    return _export_excel(rows, headers, 'Listes noires', 'listes_noires')


# ─── Détections liste noire ───────────────────────────────────────

@login_required
def liste_detections_liste_noire(request):
    query = request.GET.get('q', '').strip()
    statut = request.GET.get('statut', '')
    traite = request.GET.get('traite', '')
    items = DetectionListeNoire.objects.select_related(
        'liste_noire__type_liste_noire', 'porte_entree'
    )
    if query:
        items = items.filter(
            Q(liste_noire__nom__icontains=query) | Q(liste_noire__prenom__icontains=query) |
            Q(liste_noire__type_liste_noire__nom__icontains=query)
        )
    if statut and statut != 'TOUS':
        items = items.filter(statut=statut)
    if traite == 'OUI':
        items = items.filter(statut='TRAITE')
    elif traite == 'NON':
        items = items.exclude(statut='TRAITE')
    items = items.order_by('-created_at')
    paginator = Paginator(items, 10)
    page_obj = paginator.get_page(request.GET.get('page', 1))
    page_links = paginator.get_elided_page_range(page_obj.number, on_each_side=2, on_ends=1)
    fp = {k: v for k, v in [('q', query), ('statut', statut), ('traite', traite)] if v and v != 'TOUS'}
    return render(request, 'liste_noire/detections/liste.html', {
        'page_obj': page_obj, 'page_links': page_links, 'query': query,
        'statut_filter': statut, 'traite_filter': traite,
        'filter_params': urlencode(fp) + '&' if fp else '',
    })


@login_required
def detail_detection_liste_noire(request, pk):
    item = get_object_or_404(
        DetectionListeNoire.objects.select_related(
            'liste_noire__type_liste_noire', 'porte_entree'
        ),
        pk=pk
    )
    score = 40
    c = (item.confiance or 'MOYENNE').upper()
    if c == 'HAUTE' or c == 'ÉLEVÉE' or c == 'ELEVEE':
        score = 90
    elif c == 'MOYENNE' or c == 'MOYEN':
        score = 70
    return render(request, 'liste_noire/detections/detail.html', {
        'item': item,
        'score': score,
    })


@login_required
def traiter_detection_liste_noire(request, pk):
    item = get_object_or_404(DetectionListeNoire, pk=pk)
    if request.method == 'POST':
        item.statut = 'TRAITE'
        item.save()
        HistoriqueAction.log(request, action='MODIFICATION', entite='DetectionListeNoire', entite_id=item.pk, details=f'Détection #{item.pk} traitée')
        messages.success(request, 'Détection marquée comme traitée.')
    return redirect('liste_detections_liste_noire')


@login_required
def export_pdf_detections_liste_noire(request):
    items = DetectionListeNoire.objects.select_related(
        'liste_noire__type_liste_noire', 'porte_entree'
    ).order_by('-created_at')
    headers = ['N°', 'Visiteur', 'Liste noire', 'Porte entrée', 'Confiance', 'Statut', 'Date détection']
    rows = [[i + 1,
             f'{d.liste_noire.prenom} {d.liste_noire.nom}' if d.liste_noire else '-',
             d.liste_noire.type_liste_noire.nom if d.liste_noire and d.liste_noire.type_liste_noire else '-',
             d.porte_entree.titre if d.porte_entree else '-',
             d.confiance or '-', d.statut,
             d.date_detection.strftime('%d/%m/%Y %H:%M') if d.date_detection else '-']
            for i, d in enumerate(items)]
    return _export_pdf(rows, headers, 'Détections liste noire', 'detections_liste_noire', [0.5, 2, 2, 1.5, 1, 1, 1.5])


@login_required
def export_excel_detections_liste_noire(request):
    items = DetectionListeNoire.objects.select_related(
        'liste_noire__type_liste_noire', 'porte_entree'
    ).order_by('-created_at')
    headers = ['N°', 'Visiteur', 'Liste noire', 'Porte entrée', 'Confiance', 'Statut', 'Date détection']
    rows = [[i + 1,
             f'{d.liste_noire.prenom} {d.liste_noire.nom}' if d.liste_noire else '',
             d.liste_noire.type_liste_noire.nom if d.liste_noire and d.liste_noire.type_liste_noire else '',
             d.porte_entree.titre if d.porte_entree else '',
             d.confiance or '', d.statut,
             d.date_detection.strftime('%d/%m/%Y %H:%M') if d.date_detection else '']
            for i, d in enumerate(items)]
    return _export_excel(rows, headers, 'Détections liste noire', 'detections_liste_noire')
