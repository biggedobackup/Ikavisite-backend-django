import base64
import datetime
import json
from io import BytesIO
from urllib.parse import urlencode

from django.conf import settings
from django.contrib.auth.decorators import login_required
from django.contrib import messages
from django.core.exceptions import ValidationError
from django.core.files.base import ContentFile
from django.core.paginator import Paginator
from django.core.cache import cache
from django.db.models import Count, Q
from django.http import HttpResponse, JsonResponse
import datetime
from django.shortcuts import render, redirect, get_object_or_404
from django.utils import timezone
from channels.layers import get_channel_layer
from asgiref.sync import async_to_sync

from utilisateurs.models import HistoriqueAction
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.units import cm
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle

from alertes_et_notifications.models import Alerte
from .models import TypeVisite, Visite, Visiteur
from entreprise.models import ParametreEntreprise, PorteEntree, Personnel, Departement
from incidents.models import Incident, DetectionIncident
from liste_noire.models import ListeNoire
from objets_oublies.models import ObjetOublie


# ─── Helpers ─────────────────────────────────────────────────────────────

def _base_qs():
    return Visite.objects.all().select_related(
        'type_visite', 'visiteur', 'porte_entree',
        'personnel__departement',
    )



def _export_pdf(rows, headers, title, filename, col_ratios=None):
    buffer = BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=landscape(A4), title=title,
                            leftMargin=1.5*cm, rightMargin=1.5*cm)
    styles = getSampleStyleSheet()
    cell_style = ParagraphStyle('CellStyle', parent=styles['Normal'],
                                fontSize=8, leading=10, alignment=1)
    elements = []
    elements.append(Paragraph(title, styles['Title']))
    elements.append(Spacer(1, 0.5 * cm))
    elements.append(Paragraph(f"Généré le {timezone.now().strftime('%d/%m/%Y %H:%M')}", styles['Normal']))
    elements.append(Spacer(1, 0.5 * cm))
    col_count = len(headers)
    available_width = landscape(A4)[0] - 3 * cm
    if col_ratios is None:
        col_ratios = [1] * col_count
    total_ratio = sum(col_ratios)
    col_widths = [available_width * r / total_ratio for r in col_ratios]
    wrapped_rows = [[Paragraph(str(c), cell_style) for c in row] for row in rows]
    data = [headers] + wrapped_rows
    table = Table(data, colWidths=col_widths, repeatRows=1)
    table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1270b8')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, -1), 8),
        ('BOTTOMPADDING', (0, 0), (-1, 0), 8),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#e2e8f0')),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.HexColor('#f8fafc'), colors.white]),
        ('VALIGN', (0, 0), (-1, -1), 'TOP'),
    ]))
    elements.append(table)
    doc.build(elements)
    pdf = buffer.getvalue()
    buffer.close()
    response = HttpResponse(pdf, content_type='application/pdf')
    response['Content-Disposition'] = f'inline; filename="{filename}_{timezone.now().strftime("%Y%m%d")}.pdf"'
    return response


def _export_excel(rows, headers, title, filename):
    wb = Workbook()
    ws = wb.active
    ws.title = title[:31]
    ws.append(headers)
    for row in rows:
        ws.append(row)
    for col_idx in range(1, len(headers) + 1):
        ws.column_dimensions[chr(64 + col_idx)].width = 22
    header_fill = PatternFill(start_color='1270b8', end_color='1270b8', fill_type='solid')
    header_font = Font(bold=True, color='ffffff', size=11)
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    response = HttpResponse(buffer, content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = f'attachment; filename="{filename}_{timezone.now().strftime("%Y%m%d")}.xlsx"'
    return response


# ─── Types de visite (CRUD) ─────────────────────────────────────────────

@login_required
def liste_types_visite(request):
    query = request.GET.get('q', '').strip()
    items = TypeVisite.objects.all()
    if query:
        items = items.filter(Q(nom__icontains=query) | Q(description__icontains=query))
    items = items.order_by('nom')
    paginator = Paginator(items, 10)
    page_obj = paginator.get_page(request.GET.get('page', 1))
    page_links = paginator.get_elided_page_range(page_obj.number, on_each_side=2, on_ends=1)
    return render(request, 'visites/types/liste.html', {
        'page_obj': page_obj, 'page_links': page_links, 'query': query,
    })


@login_required
def ajouter_type_visite(request):
    if request.method == 'POST':
        nom = request.POST.get('nom', '').strip()
        if not nom:
            messages.error(request, 'Le nom est requis.')
            return render(request, 'visites/types/ajouter.html')
        if TypeVisite.objects.filter(nom=nom).exists():
            messages.error(request, 'Ce type de visite existe déjà.')
            return render(request, 'visites/types/ajouter.html')
        type_ = TypeVisite.objects.create(
            nom=nom,
            description=request.POST.get('description', '').strip() or None,
            statut=request.POST.get('statut', 'ACTIF'),
            created_by=request.user,
        )
        HistoriqueAction.log(request, 'AJOUT', 'TypeVisite', entite_id=type_.pk, details=f'Type créé : {nom}')
        messages.success(request, 'Type de visite ajouté avec succès.')
        return redirect('liste_types_visite')
    return render(request, 'visites/types/ajouter.html')


@login_required
def detail_type_visite(request, pk):
    item = get_object_or_404(TypeVisite, pk=pk)
    return render(request, 'visites/types/detail.html', {'item': item})


@login_required
def modifier_type_visite(request, pk):
    item = get_object_or_404(TypeVisite, pk=pk)
    if request.method == 'POST':
        nom = request.POST.get('nom', '').strip()
        if not nom:
            messages.error(request, 'Le nom est requis.')
            return render(request, 'visites/types/modifier.html', {'item': item})
        if TypeVisite.objects.filter(nom=nom).exclude(pk=pk).exists():
            messages.error(request, 'Ce nom existe déjà.')
            return render(request, 'visites/types/modifier.html', {'item': item})
        item.nom = nom
        item.description = request.POST.get('description', '').strip() or None
        item.statut = request.POST.get('statut', 'ACTIF')
        item.updated_by = request.user
        item.save()
        HistoriqueAction.log(request, 'MODIFICATION', 'TypeVisite', entite_id=item.pk, details=f'Type modifié : {item.nom}')
        messages.success(request, 'Type de visite modifié avec succès.')
        return redirect('liste_types_visite')
    return render(request, 'visites/types/modifier.html', {'item': item})


@login_required
def supprimer_type_visite(request, pk):
    item = get_object_or_404(TypeVisite, pk=pk)
    if request.method == 'POST':
        item.delete()
        HistoriqueAction.log(request, 'SUPPRESSION', 'TypeVisite', entite_id=item.pk, details=f'Type supprimé : {item.nom}')
        messages.success(request, 'Type de visite supprimé avec succès.')
        return redirect('liste_types_visite')
    return render(request, 'visites/types/detail.html', {'item': item, 'confirm_delete': True})


@login_required
def export_pdf_types_visite(request):
    items = TypeVisite.objects.all().order_by('nom')
    headers = ['N°', 'Nom', 'Description', 'Statut']
    rows = [[i + 1, t.nom, t.description or '-', t.statut] for i, t in enumerate(items)]
    return _export_pdf(rows, headers, 'Types de visite', 'types_visite', [0.5, 2, 4, 1])


@login_required
def export_excel_types_visite(request):
    items = TypeVisite.objects.all().order_by('nom')
    headers = ['N°', 'Nom', 'Description', 'Statut']
    rows = [[i + 1, t.nom, t.description or '', t.statut] for i, t in enumerate(items)]
    return _export_excel(rows, headers, 'Types de visite', 'types_visite')


# ─── Visites (CRUD) ─────────────────────────────────────────────────────

@login_required
def liste_visites(request):
    query = request.GET.get('q', '').strip()
    statut = request.GET.get('statut', '')
    genre = request.GET.get('genre', '')
    departement_id = request.GET.get('departement', '')
    type_visite_id = request.GET.get('type_visite', '')
    date_start = request.GET.get('date_start', '')
    date_end = request.GET.get('date_end', '')
    items = _base_qs()
    if query:
        items = items.filter(
            Q(visiteur__nom__icontains=query) | Q(visiteur__prenom__icontains=query) |
            Q(motif__icontains=query) | Q(observations__icontains=query)
        )
    if statut:
        items = items.filter(statut=statut)
    if genre:
        items = items.filter(Q(visiteur__genre=genre) | Q(genre=genre))
    if departement_id:
        items = items.filter(personnel__departement_id=departement_id)
    if type_visite_id:
        items = items.filter(type_visite_id=type_visite_id)
    if date_start:
        items = items.filter(date_visite__gte=date_start)
    if date_end:
        items = items.filter(date_visite__lte=date_end)
    items = items.order_by('-date_visite')
    paginator = Paginator(items, 10)
    page_obj = paginator.get_page(request.GET.get('page', 1))
    departements = Departement.objects.filter(statut='ACTIF')
    types_visite = TypeVisite.objects.filter(statut='ACTIF')
    page_links = paginator.get_elided_page_range(page_obj.number, on_each_side=2, on_ends=1)
    fp = {k: v for k, v in [('q', query), ('statut', statut), ('genre', genre),
                             ('departement', departement_id), ('type_visite', type_visite_id),
                             ('date_start', date_start), ('date_end', date_end)] if v}
    return render(request, 'visites/liste.html', {
        'page_obj': page_obj, 'page_links': page_links, 'query': query, 'statut_filter': statut,
        'genre_filter': genre, 'departement_filter': departement_id,
        'type_visite_filter': type_visite_id,
        'date_start_filter': date_start, 'date_end_filter': date_end,
        'departements': departements, 'types_visite': types_visite,
        'filter_params': urlencode(fp) + '&' if fp else '',
    })


NATIONALITE_OPTIONS = [
    'Afghane', 'Albanaise', 'Algérienne', 'Allemande', 'Américaine', 'Andorrane',
    'Angolaise', 'Argentine', 'Arménienne', 'Australienne', 'Autrichienne',
    'Bahamienne', 'Bahreïnienne', 'Bangladeshie', 'Barbadienne', 'Belge',
    'Bélizienne', 'Béninoise', 'Bhoutanaise', 'Biélorusse', 'Birmane',
    'Bolivienne', 'Bosnienne', 'Botswanaise', 'Brésilienne', 'Britannique',
    'Brunéienne', 'Bulgare', 'Burkinabè', 'Burundaise',
    'Cambodgienne', 'Camerounaise', 'Canadienne', 'Cap-verdienne',
    'Centrafricaine', 'Chilienne', 'Chinoise', 'Chypriote', 'Colombienne',
    'Comorienne', 'Congolaise', 'Coréenne', 'Costaricaine', 'Croate', 'Cubaine',
    'Danoise', 'Djiboutienne', 'Dominicaine', 'Dominiquaise',
    'Égyptienne', 'Émiratie', 'Équatorienne', 'Érythréenne', 'Espagnole',
    'Estonienne', 'Eswatinienne', 'Éthiopienne',
    'Fidjienne', 'Finlandaise', 'Française',
    'Gabonaise', 'Gambienne', 'Géorgienne', 'Ghanéenne', 'Grecque',
    'Grenadienne', 'Guatémaltèque', 'Guinéenne', 'Bissau-Guinéenne', 'Guyanienne',
    'Haïtienne', 'Hondurienne', 'Hongroise',
    'Indienne', 'Indonésienne', 'Irakienne', 'Iranienne', 'Irlandaise',
    'Islandaise', 'Israélienne', 'Italienne', 'Ivoirienne',
    'Jamaïcaine', 'Japonaise', 'Jordanienne',
    'Kazakhstanaise', 'Kényane', 'Kirghize', 'Kiribatienne', 'Koweïtienne',
    'Laotienne', 'Lesothane', 'Lettone', 'Libanaise', 'Libérienne', 'Libyenne',
    'Liechtensteinoise', 'Lituanienne', 'Luxembourgeoise',
    'Macédonienne', 'Malgache', 'Malaisienne', 'Malawienne', 'Maldivienne',
    'Malienne', 'Maltaise', 'Marocaine', 'Marshallaise', 'Mauritanienne',
    'Mauricienne', 'Mexicaine', 'Micronésienne', 'Moldave', 'Monégasque',
    'Mongole', 'Monténégrine', 'Mozambicaine',
    'Namibienne', 'Nauruane', 'Néerlandaise', 'Néo-zélandaise', 'Népalaise',
    'Nicaraguayenne', 'Nigérienne', 'Nigériane', 'Nord-coréenne', 'Norvégienne',
    'Omanaise', 'Ougandaise', 'Ouzbèke',
    'Pakistanaise', 'Palaosienne', 'Palestinienne', 'Panaméenne',
    'Papouane-néo-guinéenne', 'Paraguayenne', 'Péruvienne', 'Philippine',
    'Polonaise', 'Portugaise',
    'Qatarie',
    'Roumaine', 'Russe', 'Rwandaise',
    'Saint-lucienne', 'Saint-marinaise', 'Salomonaise', 'Salvadorienne',
    'Samoane', 'Santoméenne', 'Saoudienne', 'Sénégalaise', 'Serbe',
    'Seychelloise', 'Sierra-léonaise', 'Singapourienne', 'Slovaque',
    'Slovène', 'Somalienne', 'Soudanaise', 'Sri-lankaise', 'Sud-africaine',
    'Suédoise', 'Suisse', 'Surinamaise', 'Syrienne',
    'Tadjike', 'Tanzanienne', 'Tchadienne', 'Tchèque', 'Thaïlandaise',
    'Timoraise', 'Togolaise', 'Tongienne', 'Trinidadienne', 'Tunisienne',
    'Turkmène', 'Turque', 'Tuvaluane',
    'Ukrainienne', 'Uruguayenne',
    'Vanuatuane', 'Vaticane', 'Vénézuélienne', 'Vietnamienne',
    'Yéménite',
    'Zaïroise', 'Zambienne', 'Zimbabwéenne',
]

PAYS_OPTIONS = [
    'Afghanistan', 'Afrique du Sud', 'Albanie', 'Algérie', 'Allemagne', 'Angola',
    'Arabie Saoudite', 'Argentine', 'Australie', 'Autriche', 'Belgique', 'Bénin',
    'Brésil', 'Burkina Faso', 'Burundi', 'Cambodge', 'Cameroun', 'Canada',
    'Cap-Vert', 'Chine', 'Colombie', 'Comores', 'Corée du Sud', 'Costa Rica',
    "Côte d'Ivoire", 'Croatie', 'Danemark', 'Djibouti', 'Égypte', 'Émirats arabes unis',
    'Espagne', 'Estonie', 'États-Unis', 'Éthiopie', 'Finlande', 'France', 'Gabon',
    'Gambie', 'Ghana', 'Grèce', 'Guinée', 'Guinée-Bissau', 'Guinée équatoriale',
    'Haïti', 'Hongrie', 'Inde', 'Indonésie', 'Iran', 'Iraq', 'Irlande', 'Islande',
    'Israël', 'Italie', 'Japon', 'Jordanie', 'Kenya', 'Koweït', 'Laos', 'Lesotho',
    'Lettonie', 'Liban', 'Liberia', 'Libye', 'Liechtenstein', 'Lituanie', 'Luxembourg',
    'Madagascar', 'Malaisie', 'Malawi', 'Mali', 'Malte', 'Maroc', 'Maurice',
    'Mauritanie', 'Mexique', 'Monaco', 'Mongolie', 'Monténégro', 'Mozambique',
    'Namibie', 'Népal', 'Nicaragua', 'Niger', 'Nigeria', 'Norvège', 'Nouvelle-Zélande',
    'Ouganda', 'Ouzbékistan', 'Pakistan', 'Panama', 'Paraguay', 'Pays-Bas', 'Pérou',
    'Philippines', 'Pologne', 'Portugal', 'Qatar', 'République centrafricaine',
    'République démocratique du Congo', 'République du Congo', 'Roumanie',
    'Royaume-Uni', 'Russie', 'Rwanda', 'Sénégal', 'Serbie', 'Seychelles',
    'Sierra Leone', 'Singapour', 'Slovaquie', 'Slovénie', 'Somalie', 'Soudan',
    'Soudan du Sud', 'Sri Lanka', 'Suède', 'Suisse', 'Suriname', 'Syrie',
    'Tadjikistan', 'Tanzanie', 'Tchad', 'République tchèque', 'Thaïlande', 'Togo',
    'Tunisie', 'Turkménistan', 'Turquie', 'Ukraine', 'Uruguay', 'Vatican',
    'Venezuela', 'Viêt Nam', 'Yémen', 'Zambie', 'Zimbabwe',
]


def _save_signature(signature_data, prefix='signature'):
    if not signature_data:
        return None
    try:
        fmt, data = signature_data.split(';base64,')
        ext = fmt.split('/')[-1]
        return ContentFile(base64.b64decode(data), name=f'{prefix}_{timezone.now().timestamp()}.{ext}')
    except (ValueError, TypeError):
        return None


def _detect_matches(nom, prenom, numero_piece, numero_nip, save_detections=False, user=None, porte_entree_id=None):
    matches = []
    q_nom = Q(nom__iexact=nom)
    q_prenom = Q(prenom__iexact=prenom) if prenom else Q()
    q_nip = Q(numero_nip=numero_nip) if numero_nip else Q()
    q_piece = Q(numero_piece=numero_piece) if numero_piece else Q()
    q = q_nom & q_prenom
    if numero_piece:
        q |= Q(numero_piece=numero_piece)
    if numero_nip:
        q |= Q(numero_nip=numero_nip)

    for ln in ListeNoire.objects.filter(q, statut='ACTIF').select_related('type_liste_noire'):
        detail = ln.motif or (ln.type_liste_noire.nom if hasattr(ln, 'type_liste_noire') and ln.type_liste_noire else 'Inscrit en liste noire')
        matches.append({'type': 'Liste noire', 'nom': ln.nom, 'prenom': ln.prenom, 'detail': detail, 'pk': ln.pk})
        if save_detections and user and porte_entree_id:
            from liste_noire.models import DetectionListeNoire
            DetectionListeNoire.objects.create(
                liste_noire=ln,
                porte_entree_id=porte_entree_id,
                date_detection=timezone.now(),
                confiance='MOYENNE',
                statut='ACTIF',
                notes=f'Détecté lors de la création d\'une visite pour {nom} {prenom}',
                created_by=user,
            )
    for inc in Incident.objects.filter(visite__visiteur__in=Visiteur.objects.filter(q)).select_related('visite__visiteur'):
        detail = f'{inc.type_incident} — {inc.gravite}'
        matches.append({'type': 'Incident', 'nom': inc.visite.visiteur.nom if inc.visite and inc.visite.visiteur else nom, 'prenom': inc.visite.visiteur.prenom if inc.visite and inc.visite.visiteur else prenom, 'detail': detail, 'pk': inc.pk})
        if save_detections and user:
            DetectionIncident.objects.create(
                incident=inc,
                date_detection=timezone.now(),
                confiance='MOYENNE',
                statut='ACTIF',
                notes=f'Détecté lors de la création d\'une visite pour {nom} {prenom}',
                created_by=user,
            )
    for obj in ObjetOublie.objects.filter(visite__visiteur__in=Visiteur.objects.filter(q)).select_related('visite__visiteur'):
        detail = f'{obj.nom_objet} ({obj.categorie})' if obj.categorie else obj.nom_objet
        matches.append({'type': 'Objet oublié', 'nom': obj.visite.visiteur.nom if obj.visite and obj.visite.visiteur else nom, 'prenom': obj.visite.visiteur.prenom if obj.visite and obj.visite.visiteur else prenom, 'detail': detail, 'pk': obj.pk})
    if save_detections and user:
        for m in matches:
            type_map = {'Liste noire': 'LISTE_NOIRE', 'Incident': 'INCIDENT', 'Objet oublié': 'OBJET_OUBLIE'}
            alerte_type = type_map.get(m['type'])
            if alerte_type:
                Alerte.objects.create(
                    type=alerte_type,
                    entite_id=m['pk'],
                    message=f"{m['type']} — {m['prenom']} {m['nom']} — {m['detail']}",
                )
    return matches


def _base_ctx():
    entreprise = ParametreEntreprise.objects.first()
    return {
        'portes': PorteEntree.objects.filter(statut='ACTIF'),
        'types': TypeVisite.objects.filter(statut='ACTIF'),
        'personnels': Personnel.objects.filter(statut='ACTIF').select_related('departement'),
        'departements': Departement.objects.filter(statut='ACTIF'),
        'pays_options': PAYS_OPTIONS,
        'nationalites': Visiteur.objects.filter(nationalite__isnull=False).exclude(nationalite='').values_list('nationalite', flat=True).distinct().order_by('nationalite'),
        'nationalite_options': NATIONALITE_OPTIONS,
        'regula_api_url': getattr(settings, 'REGULA_API_URL', 'https://api.regulaforensics.com'),
        'regula_api_key': getattr(settings, 'REGULA_API_KEY', ''),
        'duree_moyenne_visites': entreprise.duree_moyenne_visites if entreprise and entreprise.duree_moyenne_visites else 60,
    }


@login_required
def ajouter_visite(request):
    if request.method == 'POST':
        v_prenom = request.POST.get('v_prenom', '').strip()
        v_nom = request.POST.get('v_nom', '').strip()
        if not v_prenom or not v_nom:
            messages.error(request, 'Le prénom et le nom du visiteur sont requis.')
            return render(request, 'visites/ajouter.html', _base_ctx())

        numero_piece = request.POST.get('v_numero_piece', '').strip()
        numero_nip = request.POST.get('v_nip', '').strip()

        genre_choices = {'Homme': 'Homme', 'Femme': 'Femme'}
        telephone = request.POST.get('v_telephone', '').strip() or None

        visitor_data = dict(
            nom=v_nom,
            prenom=v_prenom,
            genre=genre_choices.get(request.POST.get('v_genre', ''), request.POST.get('v_genre', '')),
            date_naissance=request.POST.get('v_date_naissance', '').strip() or None,
            lieu_naissance=request.POST.get('v_lieu_naissance', '').strip() or None,
            nationalite=request.POST.get('v_nationalite', '').strip() or None,
            profession=request.POST.get('v_profession', '').strip() or None,
            telephone=telephone,
            adresse=request.POST.get('v_adresse', '').strip() or None,
            email=request.POST.get('v_email', '').strip() or None,
            piece_identite=request.POST.get('v_piece_identite', '').strip() or None,
            numero_piece=numero_piece or None,
            numero_nip=numero_nip or None,
            pays_delivrance=request.POST.get('v_pays_delivrance', '').strip() or None,
            date_delivrance=request.POST.get('v_date_delivrance', '').strip() or None,
        )

        files = {}
        for fld in ('photo', 'document_recto', 'document_verso'):
            f = request.FILES.get(fld)
            if f:
                files[fld] = f
        visiteur = Visiteur.chercher_ou_creer(visitor_data, files)

        portrait_data = request.POST.get('photo_portrait_data', '')
        if portrait_data and not request.FILES.get('photo') and not visiteur.photo:
            visiteur.photo = _save_signature(portrait_data, 'portrait')
            visiteur.save(update_fields=['photo'])

        arrival = request.POST.get('arrival_datetime', '').strip()
        planned_departure = request.POST.get('planned_departure_datetime', '').strip()
        departure = request.POST.get('departure_datetime', '').strip()

        from datetime import datetime as dt_datetime, date as dt_date, time as dt_time
        from django.utils import timezone

        def _parse_dt_local(v):
            if not v or 'T' not in v:
                return None, None, None
            try:
                dt = dt_datetime.strptime(v, '%Y-%m-%dT%H:%M')
                dt = timezone.make_aware(dt)
                return dt, dt.date(), dt.time()
            except (ValueError, IndexError):
                return None, None, None

        arrival_dt, arrival_date, arrival_time = _parse_dt_local(arrival)
        planned_dt, planned_date, planned_time = _parse_dt_local(planned_departure)
        departure_dt, departure_date, departure_time = _parse_dt_local(departure)

        sig_entree = _save_signature(request.POST.get('signature_entree_data', ''))
        sig_sortie = _save_signature(request.POST.get('signature_sortie_data', ''))

        visite = Visite(
            type_visite_id=request.POST.get('type_visite'),
            visiteur=visiteur,
            genre=visiteur.genre,
            porte_entree_id=request.POST.get('porte_entree'),
            personnel_id=request.POST.get('personnel') or None,
            date_visite=arrival_dt,
            heure_arrivee=arrival_time,
            motif=request.POST.get('motif', '').strip() or None,
            observations=request.POST.get('observations', '').strip() or None,
            numero_badge=request.POST.get('numero_badge', '').strip() or None,
            signature_entree=sig_entree,
            signature_sortie=sig_sortie,
            date_expiration=request.POST.get('date_expiration', '').strip() or None,
            date_depart_prevue=planned_date,
            heure_depart_prevue=planned_time,
            date_depart=departure_date,
            heure_depart=departure_time,
            created_by=request.user,
        )
        try:
            visite.save()
        except ValidationError as e:
            visiteur.delete()
            for msgs in e.message_dict.values():
                for msg in msgs:
                    messages.error(request, msg)
            return render(request, 'visites/ajouter.html', _base_ctx())
        HistoriqueAction.log(request, 'AJOUT', 'Visite', entite_id=visite.pk, details=f'Visite créée pour {v_nom} {v_prenom}')
        matches = _detect_matches(v_nom, v_prenom, numero_piece, numero_nip, save_detections=True, user=request.user, porte_entree_id=visite.porte_entree_id)
        messages.success(request, 'Visite créée avec succès.')
        if matches:
            try:
                channel_layer = get_channel_layer()
                async_to_sync(channel_layer.group_send)('detections', {
                    'type': 'detection_alert',
                    'nom': v_nom, 'prenom': v_prenom, 'matches': matches,
                })
            except Exception:
                pass
            categorized = {
                'liste_noire': [m for m in matches if m['type'] == 'Liste noire'],
                'incident': [m for m in matches if m['type'] == 'Incident'],
                'objet_oublie': [m for m in matches if m['type'] == 'Objet oublié'],
            }
            ctx = _base_ctx()
            ctx['detection_matches_json'] = json.dumps({
                'nom': v_nom, 'prenom': v_prenom,
                'num_piece': numero_piece or numero_nip or '',
                'groups': categorized,
            })
            return render(request, 'visites/ajouter.html', ctx)
        return redirect('liste_visites')

    ctx = _base_ctx()
    return render(request, 'visites/ajouter.html', ctx)


@login_required
def detail_visite(request, pk):
    item = get_object_or_404(
        Visite.objects.all().select_related(
            'type_visite', 'visiteur', 'porte_entree', 'personnel__departement',
            'created_by', 'updated_by'
        ),
        pk=pk
    )
    images = []
    if item.visiteur.photo:
        images.append({'src': item.visiteur.photo.url, 'label': 'Photo visiteur', 'key': 'photo'})
    if item.visiteur.document_recto:
        images.append({'src': item.visiteur.document_recto.url, 'label': 'Document recto', 'key': 'recto'})
    if item.visiteur.document_verso:
        images.append({'src': item.visiteur.document_verso.url, 'label': 'Document verso', 'key': 'verso'})
    if item.signature_entree:
        images.append({'src': item.signature_entree.url, 'label': 'Signature entree', 'key': 'signature_entree'})
    if item.signature_sortie:
        images.append({'src': item.signature_sortie.url, 'label': 'Signature sortie', 'key': 'signature_sortie'})
    return render(request, 'visites/detail.html', {'item': item, 'visitor_images': images})


@login_required
def modifier_visite(request, pk):
    item = get_object_or_404(
        Visite.objects.all().select_related('visiteur'),
        pk=pk
    )
    if request.method == 'POST':
        v = item.visiteur
        v.prenom = request.POST.get('v_prenom', '').strip() or v.prenom
        v.nom = request.POST.get('v_nom', '').strip() or v.nom
        v.genre = request.POST.get('v_genre', '').strip() or v.genre
        v.date_naissance = request.POST.get('v_date_naissance', '').strip() or v.date_naissance
        v.lieu_naissance = request.POST.get('v_lieu_naissance', '').strip() or None
        v.nationalite = request.POST.get('v_nationalite', '').strip() or None
        v.profession = request.POST.get('v_profession', '').strip() or None
        v.telephone = request.POST.get('v_telephone', '').strip() or None
        v.adresse = request.POST.get('v_adresse', '').strip() or None
        v.email = request.POST.get('v_email', '').strip() or None
        v.piece_identite = request.POST.get('v_piece_identite', '').strip() or None
        v.numero_piece = request.POST.get('v_numero_piece', '').strip() or None
        v.numero_nip = request.POST.get('v_nip', '').strip() or None
        v.pays_delivrance = request.POST.get('v_pays_delivrance', '').strip() or None
        v.date_delivrance = request.POST.get('v_date_delivrance', '').strip() or v.date_delivrance
        portrait_data = request.POST.get('photo_portrait_data', '')
        if request.FILES.get('photo'):
            v.photo = request.FILES['photo']
        elif portrait_data:
            v.photo = _save_signature(portrait_data, 'portrait')
        if request.FILES.get('document_recto'):
            v.document_recto = request.FILES['document_recto']
        if request.FILES.get('document_verso'):
            v.document_verso = request.FILES['document_verso']
        v.save()

        from datetime import datetime as dt_datetime, date as dt_date, time as dt_time
        from django.utils import timezone

        arrival = request.POST.get('arrival_datetime', '').strip()
        planned_departure = request.POST.get('planned_departure_datetime', '').strip()
        departure = request.POST.get('departure_datetime', '').strip()

        def _parse_dt_local(v):
            if not v or 'T' not in v:
                return None, None, None
            try:
                dt = dt_datetime.strptime(v, '%Y-%m-%dT%H:%M')
                dt = timezone.make_aware(dt)
                return dt, dt.date(), dt.time()
            except (ValueError, IndexError):
                return None, None, None

        arrival_dt, arrival_date, arrival_time = _parse_dt_local(arrival)
        planned_dt, planned_date, planned_time = _parse_dt_local(planned_departure)
        departure_dt, departure_date, departure_time = _parse_dt_local(departure)

        sig_entree_data = request.POST.get('signature_entree_data', '')
        sig_sortie_data = request.POST.get('signature_sortie_data', '')

        item.type_visite_id = request.POST.get('type_visite')
        item.porte_entree_id = request.POST.get('porte_entree')
        item.personnel_id = request.POST.get('personnel') or None
        item.date_visite = arrival_dt
        item.heure_arrivee = arrival_time
        item.date_depart_prevue = planned_date
        item.heure_depart_prevue = planned_time
        item.date_depart = departure_date
        item.heure_depart = departure_time
        item.date_expiration = request.POST.get('date_expiration', '').strip() or None
        item.motif = request.POST.get('motif', '').strip() or None
        item.observations = request.POST.get('observations', '').strip() or None
        item.numero_badge = request.POST.get('numero_badge', '').strip() or None
        item.genre = v.genre
        if sig_entree_data:
            item.signature_entree = _save_signature(sig_entree_data, 'signature_entree')
        if sig_sortie_data:
            item.signature_sortie = _save_signature(sig_sortie_data, 'signature_sortie')
        item.updated_by = request.user
        item.save(skip_validation=True)
        HistoriqueAction.log(request, 'MODIFICATION', 'Visite', entite_id=item.pk, details=f'Visite modifiée — {item.visiteur.nom} {item.visiteur.prenom}')
        messages.success(request, 'Visite modifiée avec succès.')
        return redirect('liste_visites')
    ctx = _base_ctx()
    ctx['item'] = item
    return render(request, 'visites/modifier.html', ctx)


@login_required
def supprimer_visite(request, pk):
    item = get_object_or_404(Visite, pk=pk)
    if request.method == 'POST':
        HistoriqueAction.log(request, 'SUPPRESSION', 'Visite', entite_id=item.pk, details=f'Visite supprimée — {item.visiteur.nom} {item.visiteur.prenom}')
        item.delete()
        messages.success(request, 'Visite supprimée avec succès.')
        return redirect('liste_visites')
    return render(request, 'visites/detail.html', {'item': item, 'confirm_delete': True})


@login_required
def terminer_visite(request, pk):
    item = get_object_or_404(
        Visite.objects.filter(statut__in=['EN_COURS', 'EXCEDE']),
        pk=pk
    )
    if request.method == 'POST':
        now = timezone.now()
        item.statut = 'TERMINE'
        item.date_depart = now.date()
        item.heure_depart = now.time()
        item.updated_by = request.user

        sig_data = request.POST.get('signature_sortie_data', '')
        if sig_data:
            item.signature_sortie = _save_signature(sig_data, 'signature_sortie')

        incident_choice = request.POST.get('incident_choice', 'NON')
        incident_desc = request.POST.get('incident_description', '').strip()
        if incident_choice == 'OUI' and incident_desc:
            obs = item.observations or ''
            suffix = f'\nIncident: {incident_desc}'
            item.observations = (obs + suffix) if obs else suffix.lstrip()
            Incident.objects.create(
                type_incident='Incident signalé lors de la sortie',
                motif=incident_desc,
                date_incident=now,
                gravite=request.POST.get('incident_gravite', 'MOYENNE'),
                visite=item,
                created_by=request.user,
            )

        item.save(skip_validation=True)
        HistoriqueAction.log(request, 'TERMINER', 'Visite', entite_id=item.pk, details=f'Visite terminée — {item.visiteur.nom} {item.visiteur.prenom}')
        messages.success(request, 'Visite terminée avec succès.')
        return redirect('liste_visites')

    return redirect('liste_visites')


@login_required
def export_pdf_visites(request):
    items = _base_qs().order_by('-date_visite')
    headers = ['N°', 'Visiteur', 'Type', 'Porte', 'Date visite', 'Statut']
    rows = [[i + 1, str(v.visiteur), v.type_visite.nom, v.porte_entree.titre,
             v.date_visite.strftime('%d/%m/%Y %H:%M') if v.date_visite else '-', v.statut]
            for i, v in enumerate(items)]
    return _export_pdf(rows, headers, 'Liste des visites', 'visites', [0.5, 2, 1.5, 1.5, 1.5, 1])


@login_required
def export_excel_visites(request):
    items = _base_qs().order_by('-date_visite')
    headers = ['N°', 'Visiteur', 'Type', 'Porte', 'Date visite', 'Statut']
    rows = [[i + 1, str(v.visiteur), v.type_visite.nom, v.porte_entree.titre,
             v.date_visite.strftime('%d/%m/%Y %H:%M') if v.date_visite else '', v.statut]
            for i, v in enumerate(items)]
    return _export_excel(rows, headers, 'Visites', 'visites')


# ─── Visites en cours ───────────────────────────────────────────────────

def _encours_qs():
    return _base_qs().filter(statut='EN_COURS')


@login_required
def liste_visites_encours(request):
    query = request.GET.get('q', '').strip()
    genre = request.GET.get('genre', '')
    departement_id = request.GET.get('departement', '')
    type_visite_id = request.GET.get('type_visite', '')
    items = _encours_qs()
    if query:
        items = items.filter(Q(visiteur__nom__icontains=query) | Q(visiteur__prenom__icontains=query))
    if genre:
        items = items.filter(Q(visiteur__genre=genre) | Q(genre=genre))
    if departement_id:
        items = items.filter(personnel__departement_id=departement_id)
    if type_visite_id:
        items = items.filter(type_visite_id=type_visite_id)
    items = items.order_by('-date_visite')
    paginator = Paginator(items, 10)
    page_obj = paginator.get_page(request.GET.get('page', 1))
    page_links = paginator.get_elided_page_range(page_obj.number, on_each_side=2, on_ends=1)
    departements = Departement.objects.filter(statut='ACTIF')
    types_visite = TypeVisite.objects.filter(statut='ACTIF')
    fp = {k: v for k, v in [('q', query), ('genre', genre),
                             ('departement', departement_id), ('type_visite', type_visite_id)] if v}
    return render(request, 'visites/encours/liste.html', {
        'page_obj': page_obj, 'page_links': page_links, 'query': query,
        'genre_filter': genre, 'departement_filter': departement_id,
        'type_visite_filter': type_visite_id,
        'departements': departements, 'types_visite': types_visite,
        'filter_params': urlencode(fp) + '&' if fp else '',
    })


@login_required
def detail_visite_encours(request, pk):
    item = get_object_or_404(_encours_qs(), pk=pk)
    return render(request, 'visites/encours/detail.html', {'item': item})


@login_required
def export_pdf_visites_encours(request):
    items = _encours_qs().order_by('-date_visite')
    headers = ['N°', 'Visiteur', 'Type', 'Porte', 'Date visite']
    rows = [[i + 1, str(v.visiteur), v.type_visite.nom, v.porte_entree.titre,
             v.date_visite.strftime('%d/%m/%Y %H:%M') if v.date_visite else '-']
            for i, v in enumerate(items)]
    return _export_pdf(rows, headers, 'Visites en cours', 'visites_encours', [0.5, 2, 1.5, 1.5, 1.5])


@login_required
def export_excel_visites_encours(request):
    items = _encours_qs().order_by('-date_visite')
    headers = ['N°', 'Visiteur', 'Type', 'Porte', 'Date visite']
    rows = [[i + 1, str(v.visiteur), v.type_visite.nom, v.porte_entree.titre,
             v.date_visite.strftime('%d/%m/%Y %H:%M') if v.date_visite else '']
            for i, v in enumerate(items)]
    return _export_excel(rows, headers, 'Visites en cours', 'visites_encours')


# ─── Visites terminées ──────────────────────────────────────────────────

def _terminees_qs():
    return _base_qs().filter(statut='TERMINE')


@login_required
def liste_visites_terminees(request):
    query = request.GET.get('q', '').strip()
    genre = request.GET.get('genre', '')
    departement_id = request.GET.get('departement', '')
    type_visite_id = request.GET.get('type_visite', '')
    items = _terminees_qs()
    if query:
        items = items.filter(Q(visiteur__nom__icontains=query) | Q(visiteur__prenom__icontains=query))
    if genre:
        items = items.filter(Q(visiteur__genre=genre) | Q(genre=genre))
    if departement_id:
        items = items.filter(personnel__departement_id=departement_id)
    if type_visite_id:
        items = items.filter(type_visite_id=type_visite_id)
    items = items.order_by('-date_visite')
    paginator = Paginator(items, 10)
    page_obj = paginator.get_page(request.GET.get('page', 1))
    page_links = paginator.get_elided_page_range(page_obj.number, on_each_side=2, on_ends=1)
    departements = Departement.objects.filter(statut='ACTIF')
    types_visite = TypeVisite.objects.filter(statut='ACTIF')
    fp = {k: v for k, v in [('q', query), ('genre', genre),
                             ('departement', departement_id), ('type_visite', type_visite_id)] if v}
    return render(request, 'visites/terminees/liste.html', {
        'page_obj': page_obj, 'page_links': page_links, 'query': query,
        'genre_filter': genre, 'departement_filter': departement_id,
        'type_visite_filter': type_visite_id,
        'departements': departements, 'types_visite': types_visite,
        'filter_params': urlencode(fp) + '&' if fp else '',
    })


@login_required
def detail_visite_terminee(request, pk):
    item = get_object_or_404(_terminees_qs(), pk=pk)
    return render(request, 'visites/terminees/detail.html', {'item': item})


@login_required
def export_pdf_visites_terminees(request):
    items = _terminees_qs().order_by('-date_visite')
    headers = ['N°', 'Visiteur', 'Type', 'Porte', 'Date visite']
    rows = [[i + 1, str(v.visiteur), v.type_visite.nom, v.porte_entree.titre,
             v.date_visite.strftime('%d/%m/%Y %H:%M') if v.date_visite else '-']
            for i, v in enumerate(items)]
    return _export_pdf(rows, headers, 'Visites terminées', 'visites_terminees', [0.5, 2, 1.5, 1.5, 1.5])


@login_required
def export_excel_visites_terminees(request):
    items = _terminees_qs().order_by('-date_visite')
    headers = ['N°', 'Visiteur', 'Type', 'Porte', 'Date visite']
    rows = [[i + 1, str(v.visiteur), v.type_visite.nom, v.porte_entree.titre,
             v.date_visite.strftime('%d/%m/%Y %H:%M') if v.date_visite else '']
            for i, v in enumerate(items)]
    return _export_excel(rows, headers, 'Visites terminées', 'visites_terminees')


# ─── Visites excédées ───────────────────────────────────────────────────

def _excedees_qs():
    return _base_qs().filter(statut='EXCEDE')


@login_required
def liste_visites_excedees(request):
    query = request.GET.get('q', '').strip()
    genre = request.GET.get('genre', '')
    departement_id = request.GET.get('departement', '')
    type_visite_id = request.GET.get('type_visite', '')
    items = _excedees_qs()
    if query:
        items = items.filter(Q(visiteur__nom__icontains=query) | Q(visiteur__prenom__icontains=query))
    if genre:
        items = items.filter(Q(visiteur__genre=genre) | Q(genre=genre))
    if departement_id:
        items = items.filter(personnel__departement_id=departement_id)
    if type_visite_id:
        items = items.filter(type_visite_id=type_visite_id)
    items = items.order_by('-date_visite')
    paginator = Paginator(items, 10)
    page_obj = paginator.get_page(request.GET.get('page', 1))
    page_links = paginator.get_elided_page_range(page_obj.number, on_each_side=2, on_ends=1)
    departements = Departement.objects.filter(statut='ACTIF')
    types_visite = TypeVisite.objects.filter(statut='ACTIF')
    fp = {k: v for k, v in [('q', query), ('genre', genre),
                             ('departement', departement_id), ('type_visite', type_visite_id)] if v}
    return render(request, 'visites/excedees/liste.html', {
        'page_obj': page_obj, 'page_links': page_links, 'query': query,
        'genre_filter': genre, 'departement_filter': departement_id,
        'type_visite_filter': type_visite_id,
        'departements': departements, 'types_visite': types_visite,
        'filter_params': urlencode(fp) + '&' if fp else '',
    })


@login_required
def detail_visite_excedee(request, pk):
    item = get_object_or_404(_excedees_qs(), pk=pk)
    return render(request, 'visites/excedees/detail.html', {'item': item})


@login_required
def export_pdf_visites_excedees(request):
    items = _excedees_qs().order_by('-date_visite')
    headers = ['N°', 'Visiteur', 'Type', 'Porte', 'Date visite']
    rows = [[i + 1, str(v.visiteur), v.type_visite.nom, v.porte_entree.titre,
             v.date_visite.strftime('%d/%m/%Y %H:%M') if v.date_visite else '-']
            for i, v in enumerate(items)]
    return _export_pdf(rows, headers, 'Visites excédées', 'visites_excedees', [0.5, 2, 1.5, 1.5, 1.5])


@login_required
def export_excel_visites_excedees(request):
    items = _excedees_qs().order_by('-date_visite')
    headers = ['N°', 'Visiteur', 'Type', 'Porte', 'Date visite']
    rows = [[i + 1, str(v.visiteur), v.type_visite.nom, v.porte_entree.titre,
             v.date_visite.strftime('%d/%m/%Y %H:%M') if v.date_visite else '']
            for i, v in enumerate(items)]
    return _export_excel(rows, headers, 'Visites excédées', 'visites_excedees')


# ─── Visiteurs ────────────────────────────────────────────────────────────

@login_required
def liste_visiteurs(request):
    query = request.GET.get('q', '').strip()
    genre = request.GET.get('genre', '')
    nationalite = request.GET.get('nationalite', '')
    statut = request.GET.get('statut', '')
    piece_identite = request.GET.get('piece_identite', '')
    date_start = request.GET.get('date_start', '')
    date_end = request.GET.get('date_end', '')
    items = Visiteur.objects.all()
    if query:
        items = items.filter(
            Q(nom__icontains=query) | Q(prenom__icontains=query) |
            Q(email__icontains=query) | Q(telephone__icontains=query) |
            Q(nationalite__icontains=query) | Q(numero_piece__icontains=query)
        )
    if genre:
        items = items.filter(genre=genre)
    if nationalite:
        items = items.filter(nationalite=nationalite)
    if statut:
        items = items.filter(statut=statut)
    if piece_identite:
        items = items.filter(piece_identite=piece_identite)
    if date_start:
        items = items.filter(created_at__gte=date_start)
    if date_end:
        items = items.filter(created_at__lte=date_end)
    # Subquery : évite le LEFT JOIN + GROUP BY lourd de annotate(Count())
    # Ne fait la sous-requête que pour les 10 visiteurs de la page
    from django.db.models import Subquery, OuterRef, Count, IntegerField
    nb_visites_subq = Subquery(
        Visite.objects.filter(visiteur=OuterRef('pk'))
        .order_by()
        .values('visiteur')
        .annotate(count=Count('*'))
        .values('count')[:1],
        output_field=IntegerField(),
    )
    items = items.annotate(nb_visites_annotated=nb_visites_subq).order_by('-created_at')
    paginator = Paginator(items, 10)
    page_obj = paginator.get_page(request.GET.get('page', 1))
    page_links = paginator.get_elided_page_range(page_obj.number, on_each_side=2, on_ends=1)
    # Cache la liste des nationalités (peu changeante)
    cache_key = 'visiteurs_nationalites_list'
    nationalites = cache.get(cache_key)
    if nationalites is None:
        nationalites = list(Visiteur.objects.values_list('nationalite', flat=True).distinct().order_by('nationalite'))
        cache.set(cache_key, nationalites, 3600)  # 1h de cache
    fp = {k: v for k, v in [('q', query), ('genre', genre), ('nationalite', nationalite),
                             ('statut', statut), ('piece_identite', piece_identite),
                             ('date_start', date_start), ('date_end', date_end)] if v}
    return render(request, 'visiteurs/liste.html', {
        'page_obj': page_obj,
        'page_links': page_links,
        'query': query,
        'genre_filter': genre,
        'nationalite_filter': nationalite,
        'statut_filter': statut,
        'piece_identite_filter': piece_identite,
        'date_start_filter': date_start,
        'date_end_filter': date_end,
        'nationalites': nationalites,
        'filter_params': urlencode(fp) + '&' if fp else '',
    })


@login_required
def detail_visiteur(request, pk):
    item = get_object_or_404(Visiteur, pk=pk)
    visites = Visite.objects.filter(visiteur=item)\
        .select_related('type_visite', 'porte_entree', 'personnel__departement')\
        .order_by('-date_visite')
    return render(request, 'visiteurs/detail.html', {
        'item': item,
        'visites': visites,
    })


# ─── Benchmark (test uniquement) ─────────────────────────────────────────

import json
from django.views.decorators.csrf import csrf_exempt
from django.views.decorators.http import require_POST


@csrf_exempt
@require_POST
def benchmark_create_visite(request):
    import time
    start = time.time()
    try:
        data = json.loads(request.body)
        visiteur = get_object_or_404(Visiteur, pk=int(data.get('id_visiteur', 0)))

        def parse_date(v):
            if not v:
                return None
            parts = v.strip().split('T')[0].split('-')
            return datetime.date(int(parts[0]), int(parts[1]), int(parts[2]))

        def parse_time(v):
            if not v:
                return None
            parts = v.strip().split(':')
            return datetime.time(int(parts[0]), int(parts[1]))

        def parse_datetime(v):
            if not v:
                return None
            v = v.strip()
            if 'T' in v:
                d, t = v.split('T')
            else:
                d, t = v, '00:00'
            dp = d.split('-')
            tp = t.split(':')
            return datetime.datetime(int(dp[0]), int(dp[1]), int(dp[2]),
                                     int(tp[0]), int(tp[1]))

        dta = parse_datetime(data.get('date_visite'))
        if dta and timezone.is_naive(dta):
            dta = timezone.make_aware(dta)
        visite = Visite(
            type_visite_id=int(data.get('id_type_visite', 0)),
            visiteur=visiteur,
            porte_entree_id=int(data.get('id_porte_entree', 0)),
            personnel_id=data.get('id_personnel') or None,
            date_visite=dta,
            heure_arrivee=parse_time(data.get('heure_arrivee')),
            date_depart_prevue=parse_date(data.get('date_depart_prevue')),
            heure_depart_prevue=parse_time(data.get('heure_depart_prevue')),
            motif=data.get('motif', ''),
            numero_badge=data.get('numero_badge', ''),
            statut='EN_COURS',
        )
        visite.save(skip_validation=True)
        elapsed = round((time.time() - start) * 1000, 2)
        return JsonResponse({'success': True, 'id': visite.id, 'ms': elapsed})
    except Exception as e:
        elapsed = round((time.time() - start) * 1000, 2)
        return JsonResponse({'success': False, 'error': str(e), 'ms': elapsed})
